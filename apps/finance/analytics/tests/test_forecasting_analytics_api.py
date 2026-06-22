from __future__ import annotations

import csv
from contextlib import contextmanager
from datetime import date
from decimal import Decimal
from email.header import decode_header
from io import BytesIO, StringIO
from unittest.mock import patch

from django.test import override_settings
from django.urls import reverse
from django.utils.datastructures import MultiValueDict
from openpyxl import load_workbook
from rest_framework import status

from apps.finance.analytics.forecast.services import (
    ForecastHistoryPoint,
    ForecastingFilters,
    ForecastingService,
    LinearForecastModel,
    build_forecasting_filters_from_query,
)
from apps.finance.models import TransactionType
from apps.finance.testing import FinanceAPITestCase


FIXED_TODAY = date(2026, 6, 22)


@contextmanager
def patched_forecasting_today():
    with patch("apps.finance.analytics.forecast.services.get_user_app_today", return_value=FIXED_TODAY), patch(
        "apps.finance.analytics.forecast.services.timezone.localdate",
        return_value=FIXED_TODAY,
    ):
        yield


@override_settings(CURRENCY_RATES_ENABLED=False)
class ForecastingServiceTests(FinanceAPITestCase):
    def test_linear_model_builds_future_months_and_confidence_bounds(self):
        history_points = [
            ForecastHistoryPoint(month=f"2026-{month:02d}", label=f"Месяц {month}", value=Decimal(str(month * 1000)))
            for month in range(1, 7)
        ]

        result = LinearForecastModel().forecast(
            history_points=history_points,
            horizon_months=3,
            confidence=95,
            metric="expense",
        )

        self.assertEqual(result.model, "linear")
        self.assertFalse(result.used_fallback)
        self.assertEqual([point.month for point in result.forecast_points], ["2026-07", "2026-08", "2026-09"])
        self.assertEqual([point.value for point in result.forecast_points], [Decimal("7000.00"), Decimal("8000.00"), Decimal("9000.00")])
        self.assertLess(result.forecast_points[0].lower_bound, result.forecast_points[0].value)
        self.assertGreater(result.forecast_points[0].upper_bound, result.forecast_points[0].value)

    def test_build_filters_from_query_uses_defaults_and_validates_account_ownership(self):
        filters = build_forecasting_filters_from_query(
            user=self.user,
            query_params=MultiValueDict({}),
        )

        self.assertEqual(filters.metric, "expense")
        self.assertEqual(filters.source, "all")
        self.assertEqual(filters.horizon_months, 6)
        self.assertEqual(filters.confidence, 95)
        self.assertEqual(filters.model, "linear")

        with self.assertRaisesMessage(Exception, "Счёт не найден"):
            build_forecasting_filters_from_query(
                user=self.user,
                query_params=MultiValueDict({"source": [str(self.other_account.id)]}),
            )

    def test_service_returns_insufficient_data_for_less_than_six_completed_months(self):
        self.create_forecasting_dataset(months_count=5)
        filters = ForecastingFilters(
            metric="expense",
            source="all",
            horizon_months=6,
            confidence=95,
            model="linear",
        )

        with patched_forecasting_today():
            projection = ForecastingService().build_projection(user=self.user, filters=filters)

        self.assertFalse(projection["has_data"])
        self.assertTrue(projection["has_insufficient_data"])
        self.assertTrue(projection["unstable"])
        self.assertEqual(len(projection["chart_points"]), 5)
        self.assertEqual(projection["detail_rows"], [])
        self.assertIn("нужно минимум 6", projection["alert"]["text"])
        self.assertNotIn("нужноминимум", projection["alert"]["text"])

    def create_forecasting_dataset(self, *, months_count: int = 6):
        month_data = [
            (date(2025, 12, 5), "46000.00", "118000.00"),
            (date(2026, 1, 5), "78000.00", "120000.00"),
            (date(2026, 2, 5), "83000.00", "123000.00"),
            (date(2026, 3, 5), "92000.00", "125000.00"),
            (date(2026, 4, 5), "101000.00", "128000.00"),
            (date(2026, 5, 5), "112000.00", "132000.00"),
        ]
        if months_count < len(month_data):
            month_data = month_data[-months_count:]
        else:
            month_data = month_data[:months_count]
        for operation_date, expense_amount, income_amount in month_data:
            self.create_transaction(
                account=self.account,
                category=self.expense_category,
                type=TransactionType.EXPENSE,
                amount=expense_amount,
                description=f"Расход {operation_date:%Y-%m}",
                operation_date=operation_date,
            )
            self.create_transaction(
                account=self.account,
                category=self.income_category,
                type=TransactionType.INCOME,
                amount=income_amount,
                description=f"Доход {operation_date:%Y-%m}",
                operation_date=operation_date,
            )


@override_settings(CURRENCY_RATES_ENABLED=False)
class ForecastingAnalyticsAPITests(FinanceAPITestCase):
    def setUp(self):
        super().setUp()
        self.meta_url = reverse("finance:forecasting-analytics-meta")
        self.projection_url = reverse("finance:forecasting-analytics-projection")
        self.export_url = reverse("finance:forecasting-analytics-export")

    def assert_attachment_header(self, response, expected_extension: str):
        raw_header = response["Content-Disposition"]
        decoded_header = "".join(
            part.decode(encoding or "utf-8") if isinstance(part, bytes) else part
            for part, encoding in decode_header(raw_header)
        )
        self.assertIn("attachment;", decoded_header)
        self.assertIn(expected_extension, decoded_header)

    def assert_field_error(self, response, expected_field: str):
        field_errors = response.data.get("error", {}).get("field_errors", response.data)
        self.assertIn(expected_field, field_errors)

    def create_forecasting_dataset(self, *, months_count: int = 6, include_other_user: bool = True):
        month_data = [
            (date(2025, 12, 5), "46000.00", "118000.00"),
            (date(2026, 1, 5), "78000.00", "120000.00"),
            (date(2026, 2, 5), "83000.00", "123000.00"),
            (date(2026, 3, 5), "92000.00", "125000.00"),
            (date(2026, 4, 5), "101000.00", "128000.00"),
            (date(2026, 5, 5), "112000.00", "132000.00"),
        ]
        if months_count < len(month_data):
            month_data = month_data[-months_count:]
        else:
            month_data = month_data[:months_count]
        for operation_date, expense_amount, income_amount in month_data:
            self.create_transaction(
                account=self.account,
                category=self.expense_category,
                type=TransactionType.EXPENSE,
                amount=expense_amount,
                description=f"Расход {operation_date:%Y-%m}",
                operation_date=operation_date,
            )
            self.create_transaction(
                account=self.account,
                category=self.income_category,
                type=TransactionType.INCOME,
                amount=income_amount,
                description=f"Доход {operation_date:%Y-%m}",
                operation_date=operation_date,
            )

        # Текущий неполный месяц не должен попадать в историю прогноза.
        self.create_transaction(
            account=self.account,
            category=self.expense_category,
            type=TransactionType.EXPENSE,
            amount="999999.00",
            description="Июньский неполный месяц",
            operation_date=date(2026, 6, 10),
        )

        if include_other_user:
            self.create_transaction(
                user=self.other_user,
                account=self.other_account,
                category=self.other_category,
                type=TransactionType.EXPENSE,
                amount="999999.00",
                description="Чужой расход не должен попасть в прогноз",
                operation_date=date(2026, 5, 5),
            )

    def create_cash_account_forecasting_dataset(self):
        for month in range(12, 13):
            self.create_transaction(
                account=self.cash_account,
                category=self.expense_category,
                type=TransactionType.EXPENSE,
                amount="1000.00",
                description="Наличные декабрь",
                operation_date=date(2025, month, 7),
            )
        for month in range(1, 6):
            self.create_transaction(
                account=self.cash_account,
                category=self.expense_category,
                type=TransactionType.EXPENSE,
                amount=str(1000 + month * 100),
                description=f"Наличные {month}",
                operation_date=date(2026, month, 7),
            )

    def create_unstable_forecasting_dataset(self):
        month_data = [
            (date(2025, 12, 5), "1000.00"),
            (date(2026, 1, 5), "20000.00"),
            (date(2026, 2, 5), "1000.00"),
            (date(2026, 3, 5), "22000.00"),
            (date(2026, 4, 5), "1000.00"),
            (date(2026, 5, 5), "24000.00"),
        ]
        for operation_date, amount in month_data:
            self.create_transaction(
                account=self.account,
                category=self.expense_category,
                type=TransactionType.EXPENSE,
                amount=amount,
                description=f"Нестабильный расход {operation_date:%Y-%m}",
                operation_date=operation_date,
            )

    def test_forecasting_analytics_requires_authentication(self):
        meta_response = self.client.get(self.meta_url)
        projection_response = self.client.get(self.projection_url)
        export_response = self.client.post(
            self.export_url,
            data={"format": "csv", "sections": ["forecast"]},
            format="json",
        )

        self.assertEqual(meta_response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(projection_response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(export_response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_meta_returns_options_sources_and_defaults(self):
        self.authenticate()

        response = self.client.get(self.meta_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn({"value": "expense", "label": "Расходы"}, response.data["metric_options"])
        self.assertIn({"value": "linear", "label": "Линейная регрессия"}, response.data["model_options"])
        self.assertIn({"value": "prophet", "label": "Prophet"}, response.data["model_options"])
        self.assertIn({"id": "confidence", "label": "Включить интервал доверия", "icon": "mdi-shield-check"}, response.data["export_sections"])
        self.assertEqual(response.data["default_metric"], "expense")
        self.assertEqual(response.data["default_source"], "all")
        self.assertEqual(response.data["default_horizon_months"], 6)
        self.assertEqual(response.data["default_confidence"], 95)
        self.assertEqual(response.data["default_model"], "linear")
        self.assertIn("Прогноз носит оценочный характер", response.data["disclaimer"])

        source_values = {item["value"] for item in response.data["source_options"]}
        self.assertIn("all", source_values)
        self.assertIn(str(self.account.id), source_values)
        self.assertIn(str(self.cash_account.id), source_values)
        self.assertNotIn(str(self.other_account.id), source_values)

    def test_projection_returns_linear_forecast_with_history_confidence_and_summary(self):
        self.authenticate()
        self.create_forecasting_dataset()

        with patched_forecasting_today():
            response = self.client.get(
                self.projection_url,
                data={
                    "metric": "expense",
                    "source": "all",
                    "horizon_months": "6",
                    "confidence": "95",
                    "model": "linear",
                },
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_data"])
        self.assertFalse(response.data["has_insufficient_data"])
        self.assertFalse(response.data["has_high_instability"])
        self.assertEqual(response.data["assumptions"]["history_period_label"], "Дек 2025 — Май 2026")
        self.assertEqual(response.data["assumptions"]["source_label"], "Все счета")
        self.assertEqual(response.data["assumptions"]["model_label"], "Линейная регрессия")
        self.assertEqual(response.data["assumptions"]["updated_at_label"], "22 июнь 2026")
        self.assertIsNone(response.data["alert"])

        history_points = [point for point in response.data["chart_points"] if not point["is_forecast"]]
        forecast_points = [point for point in response.data["chart_points"] if point["is_forecast"]]
        self.assertEqual([point["month"] for point in history_points], ["2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05"])
        self.assertEqual([point["month"] for point in forecast_points], ["2026-06", "2026-07", "2026-08", "2026-09", "2026-10", "2026-11"])
        self.assertEqual(history_points[0]["value"], 46000.0)
        self.assertNotEqual(history_points[-1]["value"], 999999.0)
        self.assertEqual(len(response.data["detail_rows"]), 6)
        self.assertIn("lower_bound", forecast_points[0])
        self.assertIn("upper_bound", forecast_points[0])
        self.assertGreater(response.data["metric_summary"]["total_forecast_rub"], 0)
        self.assertFalse(response.data["metric_summary"]["positive_is_good"])

    def test_projection_supports_income_balance_and_account_source(self):
        self.authenticate()
        self.create_forecasting_dataset(include_other_user=True)
        self.create_cash_account_forecasting_dataset()

        with patched_forecasting_today():
            income_response = self.client.get(
                self.projection_url,
                data={"metric": "income", "source": "all", "horizon_months": "3", "confidence": "90", "model": "linear"},
            )
            balance_response = self.client.get(
                self.projection_url,
                data={"metric": "balance", "source": "all", "horizon_months": "3", "confidence": "90", "model": "linear"},
            )
            account_response = self.client.get(
                self.projection_url,
                data={"metric": "expense", "source": str(self.cash_account.id), "horizon_months": "3", "confidence": "90", "model": "linear"},
            )

        self.assertEqual(income_response.status_code, status.HTTP_200_OK)
        self.assertEqual(balance_response.status_code, status.HTTP_200_OK)
        self.assertEqual(account_response.status_code, status.HTTP_200_OK)
        self.assertTrue(income_response.data["metric_summary"]["positive_is_good"])
        self.assertTrue(balance_response.data["metric_summary"]["positive_is_good"])
        self.assertEqual(account_response.data["assumptions"]["source_label"], self.cash_account.name)
        self.assertEqual(account_response.data["chart_points"][0]["value"], 1000.0)
        self.assertLess(account_response.data["metric_summary"]["total_forecast_rub"], income_response.data["metric_summary"]["total_forecast_rub"])

    def test_projection_returns_prophet_fallback_when_optional_model_fails(self):
        self.authenticate()
        self.create_forecasting_dataset()

        with patched_forecasting_today(), patch(
            "apps.finance.analytics.forecast.services.ProphetForecastModel.forecast",
            side_effect=RuntimeError("Prophet is not installed"),
        ):
            response = self.client.get(
                self.projection_url,
                data={"metric": "expense", "source": "all", "horizon_months": "6", "confidence": "95", "model": "prophet"},
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_data"])
        self.assertIn("fallback после ошибки Prophet", response.data["assumptions"]["model_label"])
        self.assertEqual(response.data["alert"]["type"], "info")
        self.assertIn("резервная модель", response.data["alert"]["title"].lower())

    def test_projection_marks_high_instability(self):
        self.authenticate()
        self.create_unstable_forecasting_dataset()

        with patched_forecasting_today():
            response = self.client.get(
                self.projection_url,
                data={"metric": "expense", "source": "all", "horizon_months": "3", "confidence": "95", "model": "linear"},
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_data"])
        self.assertTrue(response.data["has_high_instability"])
        self.assertTrue(response.data["unstable"])
        self.assertEqual(response.data["alert"]["type"], "warning")
        self.assertEqual(response.data["alert"]["title"], "Низкая точность прогноза")

    def test_projection_returns_insufficient_data_warning(self):
        self.authenticate()
        self.create_forecasting_dataset(months_count=5)

        with patched_forecasting_today():
            response = self.client.get(
                self.projection_url,
                data={"metric": "expense", "source": "all", "horizon_months": "6", "confidence": "95", "model": "linear"},
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(response.data["has_data"])
        self.assertTrue(response.data["has_insufficient_data"])
        self.assertEqual(response.data["detail_rows"], [])
        self.assertEqual(response.data["alert"]["type"], "warning")
        self.assertIn("нужно минимум 6", response.data["alert"]["text"])
        self.assertNotIn("нужноминимум", response.data["alert"]["text"])

    def test_projection_validates_query_params(self):
        self.authenticate()
        invalid_cases = [
            ({"metric": "bad"}, "metric"),
            ({"source": "bad"}, "source"),
            ({"source": str(self.other_account.id)}, "source"),
            ({"horizon_months": "2"}, "horizon_months"),
            ({"confidence": "70"}, "confidence"),
            ({"model": "unknown"}, "model"),
        ]

        for params, expected_field in invalid_cases:
            with self.subTest(params=params):
                response = self.client.get(self.projection_url, data=params)
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assert_field_error(response, expected_field)

    def test_export_returns_csv_xlsx_and_pdf_files(self):
        self.authenticate()
        self.create_forecasting_dataset()
        payload = {
            "format": "csv",
            "sections": ["history", "forecast", "confidence", "parameters"],
            "metric": "expense",
            "source": "all",
            "horizon_months": 6,
            "confidence": 95,
            "model": "linear",
        }

        with patched_forecasting_today():
            csv_response = self.client.post(self.export_url, data=payload, format="json")

        self.assertEqual(csv_response.status_code, status.HTTP_200_OK)
        self.assertEqual(csv_response["Content-Type"], "text/csv; charset=utf-8")
        self.assert_attachment_header(csv_response, ".csv")
        csv_text = csv_response.content.decode("utf-8-sig")
        rows = list(csv.reader(StringIO(csv_text), delimiter=";"))
        self.assertEqual(rows[0], ["Прогнозирование"])
        self.assertIn(["Прогноз"], rows)
        self.assertTrue(any(row and row[0] == "Июн 2026" for row in rows))

        xlsx_payload = {**payload, "format": "xlsx"}
        with patched_forecasting_today():
            xlsx_response = self.client.post(self.export_url, data=xlsx_payload, format="json")

        self.assertEqual(xlsx_response.status_code, status.HTTP_200_OK)
        self.assert_attachment_header(xlsx_response, ".xlsx")
        workbook = load_workbook(BytesIO(xlsx_response.content))
        self.assertIn("Сводка", workbook.sheetnames)
        self.assertIn("График", workbook.sheetnames)
        self.assertIn("Параметры", workbook.sheetnames)
        self.assertGreater(len(workbook["График"]._charts), 0)

        pdf_payload = {**payload, "format": "pdf"}
        with patched_forecasting_today():
            pdf_response = self.client.post(self.export_url, data=pdf_payload, format="json")

        self.assertEqual(pdf_response.status_code, status.HTTP_200_OK)
        self.assertEqual(pdf_response["Content-Type"], "application/pdf")
        self.assert_attachment_header(pdf_response, ".pdf")
        self.assertTrue(pdf_response.content.startswith(b"%PDF"))
        self.assertGreater(len(pdf_response.content), 1000)

    def test_export_validates_body_params(self):
        self.authenticate()
        invalid_cases = [
            ({"format": "docx", "sections": ["forecast"]}, "format"),
            ({"format": "csv", "sections": []}, "sections"),
            ({"format": "csv", "sections": ["bad"]}, "sections"),
            ({"format": "csv", "sections": ["forecast"], "metric": "bad"}, "metric"),
            ({"format": "csv", "sections": ["forecast"], "horizon_months": 2}, "horizon_months"),
            ({"format": "csv", "sections": ["forecast"], "confidence": 70}, "confidence"),
            ({"format": "csv", "sections": ["forecast"], "model": "bad"}, "model"),
        ]

        for payload, expected_field in invalid_cases:
            with self.subTest(payload=payload):
                response = self.client.post(self.export_url, data=payload, format="json")
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assert_field_error(response, expected_field)
