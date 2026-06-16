from __future__ import annotations

import csv
from email.header import decode_header
from datetime import date
from io import BytesIO, StringIO

from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status

from apps.finance.models import TransactionType
from apps.finance.testing import FinanceAPITestCase


@override_settings(CURRENCY_RATES_ENABLED=False)
class CombinedAnalyticsAPITests(FinanceAPITestCase):
    def setUp(self):
        super().setUp()
        self.meta_url = reverse("finance:combined-analytics-meta")
        self.aggregates_url = reverse("finance:combined-analytics-aggregates")
        self.export_url = reverse("finance:combined-analytics-export")

        self.expense_category.color = "#22C55E"
        self.expense_category.save(update_fields=["color", "updated_at"])
        self.transport_category.color = "#3B82F6"
        self.transport_category.save(update_fields=["color", "updated_at"])
        self.income_category.color = "#10B981"
        self.income_category.save(update_fields=["color", "updated_at"])

    def assert_attachment_header(self, response, expected_extension: str):
        raw_header = response["Content-Disposition"]
        decoded_header = "".join(
            part.decode(encoding or "utf-8") if isinstance(part, bytes) else part
            for part, encoding in decode_header(raw_header)
        )
        self.assertIn("attachment;", decoded_header)
        self.assertIn(expected_extension, decoded_header)

    def create_combined_analytics_dataset(self):
        self.create_transaction(
            account=self.account,
            category=self.expense_category,
            type=TransactionType.EXPENSE,
            amount="4000.00",
            description="Июньские продукты",
            operation_date=date(2026, 6, 3),
        )
        self.create_transaction(
            account=self.cash_account,
            category=self.transport_category,
            type=TransactionType.EXPENSE,
            amount="2000.00",
            description="Июньский транспорт",
            operation_date=date(2026, 6, 5),
        )
        self.create_transaction(
            account=self.account,
            category=self.expense_category,
            type=TransactionType.EXPENSE,
            amount="3000.00",
            description="Майские продукты",
            operation_date=date(2026, 5, 4),
        )
        self.create_transaction(
            account=self.account,
            category=self.income_category,
            type=TransactionType.INCOME,
            amount="10000.00",
            description="Июньская зарплата",
            operation_date=date(2026, 6, 1),
        )
        self.create_transaction(
            user=self.other_user,
            account=self.other_account,
            category=self.other_category,
            type=TransactionType.EXPENSE,
            amount="999999.00",
            description="Чужой расход не должен попасть в аналитику",
            operation_date=date(2026, 6, 3),
        )

    def get_category_payload(self, response, category_id):
        items = {
            item["category_id"]: item
            for item in response.data["pie_slices"]
        }
        return items[str(category_id)]

    def test_combined_analytics_requires_authentication(self):
        meta_response = self.client.get(self.meta_url)
        aggregates_response = self.client.get(self.aggregates_url)
        export_response = self.client.post(self.export_url, data={"format": "csv", "sections": ["table"]})

        self.assertEqual(meta_response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(aggregates_response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(export_response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_meta_returns_filter_options_and_defaults_for_current_user(self):
        self.authenticate()

        response = self.client.get(self.meta_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn({"value": "month", "label": "Месяц"}, response.data["period_presets"])
        self.assertIn({"value": "expense", "label": "Расходы"}, response.data["type_options"])
        self.assertIn({"id": "pie", "label": "Включить круговую диаграмму", "icon": "mdi-chart-donut"}, response.data["export_sections"])
        self.assertEqual(response.data["default_operation_type"], "expense")
        self.assertEqual(response.data["default_account_ids"], [])
        self.assertEqual(response.data["default_category_ids"], [])

        account_values = {item["value"] for item in response.data["account_options"]}
        category_values = {item["value"] for item in response.data["category_options"]}
        filter_account_ids = {item["id"] for item in response.data["filter_accounts"]}
        filter_category_ids = {item["id"] for item in response.data["filter_categories"]}

        self.assertIn("all", account_values)
        self.assertIn(str(self.account.id), account_values)
        self.assertIn(str(self.cash_account.id), account_values)
        self.assertNotIn(str(self.other_account.id), account_values)
        self.assertIn("all", category_values)
        self.assertIn(str(self.expense_category.id), category_values)
        self.assertIn(str(self.transport_category.id), category_values)
        self.assertNotIn(str(self.other_category.id), category_values)
        self.assertIn(str(self.account.id), filter_account_ids)
        self.assertIn(str(self.expense_category.id), filter_category_ids)

    def test_aggregates_return_pie_bar_line_and_table_for_month(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "operation_type": "expense",
                "currency": "RUB",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_data"])
        self.assertEqual(response.data["period_label"], "Июнь 2026")
        self.assertEqual(response.data["total_expense_rub"], 6000.0)
        self.assertEqual(response.data["total_income_rub"], 0.0)

        groceries = self.get_category_payload(response, self.expense_category.id)
        transport = self.get_category_payload(response, self.transport_category.id)
        self.assertEqual(groceries["category_name"], "Продукты")
        self.assertEqual(groceries["category_color"], "#22C55E")
        self.assertEqual(groceries["amount_rub"], 4000.0)
        self.assertEqual(groceries["percent"], 66.67)
        self.assertEqual(transport["amount_rub"], 2000.0)
        self.assertEqual(transport["percent"], 33.33)
        aggregate_rows = {
            item["category_id"]: item
            for item in response.data["aggregate_rows"]
        }
        self.assertEqual(aggregate_rows[str(self.expense_category.id)]["amount_rub"], 4000.0)
        self.assertEqual(aggregate_rows[str(self.transport_category.id)]["percent"], 33.33)

        bar_groups = {
            item["category_id"]: item
            for item in response.data["bar_groups"]
        }
        self.assertEqual(response.data["bar_legend"]["previous_period_label"], "Май 2026")
        self.assertEqual(response.data["bar_legend"]["current_period_label"], "Июнь 2026")
        self.assertEqual(bar_groups[str(self.expense_category.id)]["previous_period_amount_rub"], 3000.0)
        self.assertEqual(bar_groups[str(self.expense_category.id)]["current_period_amount_rub"], 4000.0)
        self.assertEqual(bar_groups[str(self.transport_category.id)]["previous_period_amount_rub"], 0.0)
        self.assertEqual(bar_groups[str(self.transport_category.id)]["current_period_amount_rub"], 2000.0)

        line_points = {item["month"]: item for item in response.data["line_points"]}
        self.assertEqual(list(line_points), [
            "2026-01",
            "2026-02",
            "2026-03",
            "2026-04",
            "2026-05",
            "2026-06",
        ])
        self.assertEqual(line_points["2026-05"]["expense_rub"], 3000.0)
        self.assertEqual(line_points["2026-06"]["expense_rub"], 6000.0)
        self.assertEqual(line_points["2026-06"]["income_rub"], 0.0)
        self.assertEqual(line_points["2026-06"]["balance_rub"], -6000.0)

    def test_aggregates_filter_by_account_category_and_income_type(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        account_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "account_ids": str(self.account.id),
                "operation_type": "expense",
            },
        )
        category_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "category_ids": str(self.transport_category.id),
                "operation_type": "expense",
            },
        )
        income_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "operation_type": "income",
            },
        )

        self.assertEqual(account_response.status_code, status.HTTP_200_OK)
        self.assertEqual(account_response.data["total_expense_rub"], 4000.0)
        self.assertEqual(len(account_response.data["pie_slices"]), 1)
        self.assertEqual(account_response.data["pie_slices"][0]["category_id"], str(self.expense_category.id))

        self.assertEqual(category_response.status_code, status.HTTP_200_OK)
        self.assertEqual(category_response.data["total_expense_rub"], 2000.0)
        self.assertEqual(category_response.data["pie_slices"][0]["category_id"], str(self.transport_category.id))

        self.assertEqual(income_response.status_code, status.HTTP_200_OK)
        self.assertEqual(income_response.data["total_income_rub"], 10000.0)
        self.assertEqual(income_response.data["total_expense_rub"], 0.0)
        self.assertEqual(income_response.data["pie_slices"][0]["category_id"], str(self.income_category.id))

    def test_aggregates_accept_repeated_and_comma_separated_ids(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        repeated_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "account_ids": [str(self.account.id), str(self.cash_account.id)],
                "operation_type": "expense",
            },
        )
        comma_response = self.client.get(
            f"{self.aggregates_url}?period_preset=month&period_id=2026-06"
            f"&account_ids={self.account.id},{self.cash_account.id}&operation_type=expense"
        )

        self.assertEqual(repeated_response.status_code, status.HTTP_200_OK)
        self.assertEqual(comma_response.status_code, status.HTTP_200_OK)
        self.assertEqual(repeated_response.data["total_expense_rub"], 6000.0)
        self.assertEqual(comma_response.data["total_expense_rub"], 6000.0)

    def test_aggregates_return_empty_payload_for_period_without_data(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "custom",
                "date_from": "2025-01-01",
                "date_to": "2025-01-31",
                "operation_type": "expense",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(response.data["has_data"])
        self.assertEqual(response.data["pie_slices"], [])
        self.assertEqual(response.data["bar_groups"], [])
        self.assertEqual(response.data["aggregate_rows"], [])
        self.assertEqual(response.data["total_expense_rub"], 0.0)
        self.assertEqual(response.data["period_label"], "01.01.2025 - 31.01.2025")

    def test_aggregates_validate_filters(self):
        self.authenticate()

        invalid_period_response = self.client.get(
            self.aggregates_url,
            data={"period_preset": "wrong"},
        )
        invalid_date_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "custom",
                "date_from": "2026-06-10",
                "date_to": "2026-06-01",
            },
        )
        invalid_account_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "account_ids": "999999",
            },
        )
        invalid_category_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "category_ids": "abc",
            },
        )
        invalid_currency_response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "currency": "RUBLE",
            },
        )

        self.assertEqual(invalid_period_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_date_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_account_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_category_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_currency_response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_transfer_operation_type_returns_empty_aggregates(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        response = self.client.get(
            self.aggregates_url,
            data={
                "period_preset": "month",
                "period_id": "2026-06",
                "operation_type": "transfer",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(response.data["has_data"])
        self.assertEqual(response.data["total_expense_rub"], 0.0)
        self.assertEqual(response.data["total_income_rub"], 0.0)
        self.assertEqual(response.data["pie_slices"], [])
        self.assertEqual(response.data["bar_groups"], [])
        self.assertTrue(response.data["line_points"])
        self.assertTrue(all(point["balance_rub"] == 0.0 for point in response.data["line_points"]))

    def test_export_csv_returns_excel_friendly_file(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        response = self.client.post(
            self.export_url,
            data={
                "format": "csv",
                "sections": ["pie", "bar", "line", "table"],
                "period_preset": "month",
                "period_id": "2026-06",
                "operation_type": "expense",
                "currency": "RUB",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("text/csv", response["Content-Type"])
        self.assert_attachment_header(response, ".csv")
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))

        decoded_content = response.content.decode("utf-8-sig")
        self.assertIn("Комбинированная аналитика", decoded_content)
        self.assertIn("Круговая диаграмма", decoded_content)
        self.assertIn("Столбчатая диаграмма", decoded_content)
        self.assertIn("Линейный график", decoded_content)
        self.assertIn("Таблица агрегатов", decoded_content)

        rows = list(csv.reader(StringIO(decoded_content), delimiter=";"))
        self.assertIn(["Категория", "Сумма", "Доля, %"], rows)
        self.assertIn(["Продукты", "4000.0", "66.67"], rows)

    def test_export_xlsx_returns_workbook_with_charts(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        response = self.client.post(
            self.export_url,
            data={
                "format": "xlsx",
                "sections": ["pie", "bar", "line", "table"],
                "period_preset": "month",
                "period_id": "2026-06",
                "operation_type": "expense",
                "currency": "RUB",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response["Content-Type"],
        )
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Сводка",
                "Круговая диаграмма",
                "Столбчатая диаграмма",
                "Линейный график",
                "Таблица агрегатов",
            ],
        )
        self.assertEqual(workbook["Сводка"]["A1"].value, "Комбинированная аналитика")
        self.assertEqual(workbook["Круговая диаграмма"]["A1"].value, "Категория")
        self.assertEqual(workbook["Круговая диаграмма"]["A2"].value, "Продукты")
        self.assertTrue(workbook["Круговая диаграмма"]._charts)
        self.assertTrue(workbook["Столбчатая диаграмма"]._charts)
        self.assertTrue(workbook["Линейный график"]._charts)
        self.assertFalse(workbook["Таблица агрегатов"]._charts)

    def test_export_pdf_returns_binary_report(self):
        self.authenticate()
        self.create_combined_analytics_dataset()

        response = self.client.post(
            self.export_url,
            data={
                "format": "pdf",
                "sections": ["pie", "bar", "line", "table"],
                "period_preset": "month",
                "period_id": "2026-06",
                "operation_type": "expense",
                "currency": "RUB",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("application/pdf", response["Content-Type"])
        self.assert_attachment_header(response, ".pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertGreater(len(response.content), 1000)

    def test_export_validates_format_sections_and_filters(self):
        self.authenticate()

        invalid_format_response = self.client.post(
            self.export_url,
            data={"format": "docx", "sections": ["table"]},
            format="json",
        )
        empty_sections_response = self.client.post(
            self.export_url,
            data={"format": "csv", "sections": []},
            format="json",
        )
        invalid_section_response = self.client.post(
            self.export_url,
            data={"format": "csv", "sections": ["wrong"]},
            format="json",
        )
        invalid_filter_response = self.client.post(
            self.export_url,
            data={
                "format": "csv",
                "sections": ["table"],
                "period_preset": "custom",
                "date_from": "2026-06-10",
                "date_to": "2026-06-01",
            },
            format="json",
        )

        self.assertEqual(invalid_format_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(empty_sections_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_section_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_filter_response.status_code, status.HTTP_400_BAD_REQUEST)
