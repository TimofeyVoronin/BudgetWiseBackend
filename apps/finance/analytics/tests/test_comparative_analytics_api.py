from __future__ import annotations

import csv
from datetime import date
from email.header import decode_header
from io import BytesIO, StringIO

from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status

from apps.finance.models import Account, TransactionType
from apps.finance.testing import FinanceAPITestCase


@override_settings(CURRENCY_RATES_ENABLED=False)
class ComparativeAnalyticsAPITests(FinanceAPITestCase):
    def setUp(self):
        super().setUp()
        self.meta_url = reverse("finance:comparative-analytics-meta")
        self.comparison_url = reverse("finance:comparative-analytics-comparison")
        self.export_url = reverse("finance:comparative-analytics-export")

        self.account.color = "#4F46E5"
        self.account.save(update_fields=["color", "updated_at"])
        self.cash_account.color = "#0EA5E9"
        self.cash_account.save(update_fields=["color", "updated_at"])
        self.expense_category.color = "#22C55E"
        self.expense_category.save(update_fields=["color", "updated_at"])
        self.transport_category.color = "#3B82F6"
        self.transport_category.save(update_fields=["color", "updated_at"])
        self.income_category.color = "#10B981"
        self.income_category.save(update_fields=["color", "updated_at"])

    def assert_attachment_header(self, response, expected_extension: str):
        raw_header = response["Content-Disposition"]
        decoded_header = "".join(
            part.decode(encoding or "utf-8") if isinstance(part, bytes) else part
            for part, encoding in decode_header(raw_header)
        )
        self.assertIn("attachment;", decoded_header)
        self.assertIn(expected_extension, decoded_header)

    def create_comparative_analytics_dataset(self):
        # Апрель 2026
        self.create_transaction(
            account=self.account,
            category=self.income_category,
            type=TransactionType.INCOME,
            amount="10000.00",
            description="Апрельская зарплата",
            operation_date=date(2026, 4, 1),
        )
        self.create_transaction(
            account=self.account,
            category=self.expense_category,
            type=TransactionType.EXPENSE,
            amount="2000.00",
            description="Апрельские продукты",
            operation_date=date(2026, 4, 3),
        )
        self.create_transaction(
            account=self.account,
            category=self.transport_category,
            type=TransactionType.EXPENSE,
            amount="1000.00",
            description="Апрельский транспорт",
            operation_date=date(2026, 4, 4),
        )

        # Май 2026
        self.create_transaction(
            account=self.account,
            category=self.income_category,
            type=TransactionType.INCOME,
            amount="12000.00",
            description="Майская зарплата",
            operation_date=date(2026, 5, 1),
        )
        self.create_transaction(
            account=self.account,
            category=self.expense_category,
            type=TransactionType.EXPENSE,
            amount="4000.00",
            description="Майские продукты",
            operation_date=date(2026, 5, 3),
        )
        self.create_transaction(
            account=self.account,
            category=self.transport_category,
            type=TransactionType.EXPENSE,
            amount="1500.00",
            description="Майский транспорт",
            operation_date=date(2026, 5, 4),
        )

        # Июнь 2026
        self.create_transaction(
            account=self.account,
            category=self.income_category,
            type=TransactionType.INCOME,
            amount="15000.00",
            description="Июньская зарплата",
            operation_date=date(2026, 6, 1),
        )
        self.create_transaction(
            account=self.account,
            category=self.expense_category,
            type=TransactionType.EXPENSE,
            amount="6000.00",
            description="Июньские продукты",
            operation_date=date(2026, 6, 3),
        )
        self.create_transaction(
            account=self.account,
            category=self.transport_category,
            type=TransactionType.EXPENSE,
            amount="2500.00",
            description="Июньский транспорт",
            operation_date=date(2026, 6, 4),
        )

        # Чужие данные не должны попадать в сравнительную аналитику пользователя.
        self.create_transaction(
            user=self.other_user,
            account=self.other_account,
            category=self.other_category,
            type=TransactionType.INCOME,
            amount="999999.00",
            description="Чужой доход",
            operation_date=date(2026, 6, 1),
        )

    def get_group_values(self, response, group_key: str):
        groups = {group["key"]: group for group in response.data["bar_groups"]}
        return groups[group_key]["values"]

    def test_comparative_analytics_requires_authentication(self):
        meta_response = self.client.get(self.meta_url)
        comparison_response = self.client.get(self.comparison_url)
        export_response = self.client.post(
            self.export_url,
            data={"format": "csv", "sections": ["legend"]},
            format="json",
        )

        self.assertEqual(meta_response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(comparison_response.status_code, status.HTTP_401_UNAUTHORIZED)
        self.assertEqual(export_response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_meta_returns_filter_options_available_entities_and_defaults(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.get(self.meta_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["max_entities"], 5)
        self.assertIn({"value": "periods", "label": "Периоды"}, response.data["comparison_type_options"])
        self.assertIn({"value": "income", "label": "Сумма доходов"}, response.data["metric_options"])
        self.assertIn({"id": "charts", "label": "Включить графики", "icon": "mdi-chart-bar"}, response.data["export_sections"])
        self.assertEqual(response.data["default_comparison_type"], "periods")
        self.assertEqual(response.data["default_metric"], "income")
        self.assertIn("periods", response.data["available_entities"])
        self.assertIn("categories", response.data["available_entities"])
        self.assertIn("accounts", response.data["available_entities"])

        period_ids = {entity["id"] for entity in response.data["available_entities"]["periods"]}
        category_ids = {entity["id"] for entity in response.data["available_entities"]["categories"]}
        account_ids = {entity["id"] for entity in response.data["available_entities"]["accounts"]}

        self.assertIn("2026-06", period_ids)
        self.assertIn(str(self.expense_category.id), category_ids)
        self.assertIn(str(self.income_category.id), category_ids)
        self.assertIn(str(self.account.id), account_ids)
        self.assertIn(str(self.cash_account.id), account_ids)
        self.assertNotIn(str(self.other_category.id), category_ids)
        self.assertNotIn(str(self.other_account.id), account_ids)

    def test_periods_comparison_preserves_drag_and_drop_order_and_aggregates(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "periods",
                "year": "2026",
                "metric": "income",
                "entity_ids": ["2026-06", "2026-05", "2026-04"],
                "convert_to_rub": "true",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_data"])
        self.assertFalse(response.data["has_mixed_currencies"])
        self.assertEqual([entity["id"] for entity in response.data["entities"]], ["2026-06", "2026-05", "2026-04"])
        self.assertEqual([entity["amount_rub"] for entity in response.data["entities"]], [15000.0, 12000.0, 10000.0])
        self.assertEqual(self.get_group_values(response, "income"), [15000.0, 12000.0, 10000.0])
        self.assertEqual(self.get_group_values(response, "expense"), [8500.0, 5500.0, 3000.0])
        self.assertEqual(self.get_group_values(response, "balance"), [6500.0, 6500.0, 7000.0])
        self.assertEqual([item["id"] for item in response.data["bar_legend"]], ["2026-06", "2026-05", "2026-04"])

        cards = {card["label"]: card for card in response.data["summary_cards"]}
        self.assertEqual(cards["Доходы"]["amount_rub"], 37000.0)
        self.assertEqual(cards["Доходы"]["delta_percent"], 25.0)
        self.assertEqual(cards["Расходы"]["amount_rub"], 17000.0)
        self.assertFalse(cards["Расходы"]["positive_is_good"])
        self.assertEqual(cards["Сальдо"]["amount_rub"], 20000.0)

        difference_rows = {row["id"]: row for row in response.data["difference_rows"]}
        self.assertEqual(difference_rows["income"]["label_a"], "Июнь 2026")
        self.assertEqual(difference_rows["income"]["label_b"], "Май 2026")
        self.assertEqual(difference_rows["income"]["delta_rub"], 3000.0)
        self.assertEqual(difference_rows["expense"]["delta_rub"], 3000.0)
        self.assertEqual(difference_rows["balance"]["delta_rub"], 0.0)

        line_points = {item["month"]: item for item in response.data["line_points"]}
        self.assertEqual(line_points["2026-04"]["income_rub"], 10000.0)
        self.assertEqual(line_points["2026-05"]["expense_rub"], 5500.0)
        self.assertEqual(line_points["2026-06"]["income_rub"], 15000.0)
        self.assertEqual(line_points["2026-06"]["expense_rub"], 8500.0)

        category_rows = {row["name"]: row for row in response.data["category_rows"]}
        self.assertEqual(category_rows["Зарплата"]["previous_amount_rub"], 12000.0)
        self.assertEqual(category_rows["Зарплата"]["current_amount_rub"], 15000.0)
        self.assertEqual(category_rows["Зарплата"]["delta_percent"], 25.0)

    def test_categories_comparison_returns_category_entities_and_differences(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "categories",
                "year": "2026",
                "metric": "expense",
                "entity_ids": [str(self.expense_category.id), str(self.transport_category.id)],
                "convert_to_rub": "true",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_data"])
        self.assertEqual(
            [entity["id"] for entity in response.data["entities"]],
            [str(self.expense_category.id), str(self.transport_category.id)],
        )
        self.assertEqual([entity["amount_rub"] for entity in response.data["entities"]], [12000.0, 5000.0])
        self.assertEqual(self.get_group_values(response, "income"), [0.0, 0.0])
        self.assertEqual(self.get_group_values(response, "expense"), [12000.0, 5000.0])
        self.assertEqual(self.get_group_values(response, "balance"), [-12000.0, -5000.0])

        difference_rows = {row["id"]: row for row in response.data["difference_rows"]}
        self.assertEqual(difference_rows["expense"]["value_a"], 12000.0)
        self.assertEqual(difference_rows["expense"]["value_b"], 5000.0)
        self.assertEqual(difference_rows["expense"]["delta_rub"], 7000.0)
        self.assertEqual(difference_rows["expense"]["delta_percent"], 140.0)
        self.assertNotIn("999999", str(response.data))

        line_points = {item["month"]: item for item in response.data["line_points"]}
        self.assertEqual(line_points["2026-04"]["expense_rub"], 3000.0)
        self.assertEqual(line_points["2026-05"]["expense_rub"], 5500.0)
        self.assertEqual(line_points["2026-06"]["expense_rub"], 8500.0)

    def test_accounts_comparison_detects_mixed_currencies_without_forcing_conversion(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()
        usd_account = Account.objects.create(
            user=self.user,
            name="USD счёт",
            initial_balance="0.00",
            balance="0.00",
            currency="USD",
            color="#EF4444",
        )
        self.create_transaction(
            account=usd_account,
            category=self.income_category,
            type=TransactionType.INCOME,
            amount="100.00",
            description="USD income",
            operation_date=date(2026, 6, 10),
        )

        response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "accounts",
                "year": "2026",
                "metric": "income",
                "entity_ids": [str(usd_account.id), str(self.account.id)],
                "convert_to_rub": "false",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["has_mixed_currencies"])
        self.assertEqual([entity["id"] for entity in response.data["entities"]], [str(usd_account.id), str(self.account.id)])
        self.assertEqual([entity["currency_code"] for entity in response.data["entities"]], ["USD", "RUB"])
        self.assertEqual(self.get_group_values(response, "income"), [100.0, 37000.0])
        self.assertEqual(self.get_group_values(response, "expense"), [0.0, 17000.0])

    def test_comparison_uses_default_entities_when_entity_ids_are_not_passed(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "periods",
                "year": "2026",
                "metric": "income",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertLessEqual(len(response.data["entities"]), 5)
        self.assertTrue(response.data["entities"])
        self.assertEqual(response.data["entities"][0]["id"].split("-")[0], "2026")

    def test_comparison_validates_query_parameters(self):
        self.authenticate()

        invalid_type_response = self.client.get(
            self.comparison_url,
            data={"comparison_type": "wrong"},
        )
        invalid_metric_response = self.client.get(
            self.comparison_url,
            data={"comparison_type": "periods", "metric": "wrong"},
        )
        invalid_year_response = self.client.get(
            self.comparison_url,
            data={"comparison_type": "periods", "year": "abc"},
        )
        invalid_period_response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "periods",
                "year": "2026",
                "entity_ids": "2025-12",
            },
        )
        too_many_entities_response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "periods",
                "year": "2026",
                "entity_ids": ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"],
            },
        )
        invalid_category_id_response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "categories",
                "year": "2026",
                "entity_ids": "abc",
            },
        )
        missing_category_response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "categories",
                "year": "2026",
                "entity_ids": "999999",
            },
        )
        invalid_bool_response = self.client.get(
            self.comparison_url,
            data={
                "comparison_type": "accounts",
                "year": "2026",
                "convert_to_rub": "maybe",
            },
        )

        self.assertEqual(invalid_type_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_metric_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_year_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_period_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(too_many_entities_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_category_id_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(missing_category_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_bool_response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_export_csv_returns_excel_friendly_comparative_report(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.post(
            self.export_url,
            data={
                "format": "csv",
                "sections": ["legend", "differences", "charts"],
                "comparison_type": "periods",
                "year": "2026",
                "metric": "income",
                "entity_ids": ["2026-06", "2026-05", "2026-04"],
                "convert_to_rub": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("text/csv", response["Content-Type"])
        self.assert_attachment_header(response, ".csv")
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))

        decoded_content = response.content.decode("utf-8-sig")
        self.assertIn("Сравнительная аналитика", decoded_content)
        self.assertIn("Легенда сравнения", decoded_content)
        self.assertIn("Динамика показателя", decoded_content)
        self.assertIn("Различия", decoded_content)

        rows = list(csv.reader(StringIO(decoded_content), delimiter=";"))
        self.assertIn(["Показатель", "Июнь 2026", "Май 2026", "Апрель 2026"], rows)
        self.assertIn(["Доходы", "15000.0", "12000.0", "10000.0"], rows)
        self.assertIn(["Расходы", "8500.0", "5500.0", "3000.0"], rows)

    def test_export_xlsx_returns_workbook_with_tables_and_non_overlapping_charts(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.post(
            self.export_url,
            data={
                "format": "xlsx",
                "sections": ["legend", "differences", "charts"],
                "comparison_type": "periods",
                "year": "2026",
                "metric": "income",
                "entity_ids": ["2026-06", "2026-05", "2026-04"],
                "convert_to_rub": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response["Content-Type"],
        )
        self.assert_attachment_header(response, ".xlsx")

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Сводка", "Легенда", "Графики", "Различия"])
        self.assertEqual(workbook["Сводка"]["A1"].value, "Сравнительная аналитика")
        self.assertEqual(workbook["Легенда"]["A1"].value, "ID")
        self.assertEqual(workbook["Графики"]["A1"].value, "Показатель")
        self.assertEqual(workbook["Различия"]["A1"].value, "Различия")

        charts = workbook["Графики"]._charts
        self.assertEqual(len(charts), 2)
        self.assertLess(charts[0].anchor._from.row, charts[1].anchor._from.row)
        self.assertGreaterEqual(charts[1].anchor._from.row, 27)

    def test_export_pdf_returns_binary_comparative_report(self):
        self.authenticate()
        self.create_comparative_analytics_dataset()

        response = self.client.post(
            self.export_url,
            data={
                "format": "pdf",
                "sections": ["legend", "differences", "charts"],
                "comparison_type": "periods",
                "year": "2026",
                "metric": "income",
                "entity_ids": ["2026-06", "2026-05", "2026-04"],
                "convert_to_rub": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("application/pdf", response["Content-Type"])
        self.assert_attachment_header(response, ".pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertGreater(len(response.content), 1000)

    def test_export_validates_format_sections_and_filters(self):
        self.authenticate()

        invalid_format_response = self.client.post(
            self.export_url,
            data={"format": "docx", "sections": ["legend"]},
            format="json",
        )
        empty_sections_response = self.client.post(
            self.export_url,
            data={"format": "csv", "sections": []},
            format="json",
        )
        invalid_section_response = self.client.post(
            self.export_url,
            data={"format": "csv", "sections": ["wrong"]},
            format="json",
        )
        invalid_filter_response = self.client.post(
            self.export_url,
            data={
                "format": "csv",
                "sections": ["legend"],
                "comparison_type": "periods",
                "year": "2026",
                "entity_ids": ["2025-12"],
            },
            format="json",
        )

        self.assertEqual(invalid_format_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(empty_sections_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_section_response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(invalid_filter_response.status_code, status.HTTP_400_BAD_REQUEST)
