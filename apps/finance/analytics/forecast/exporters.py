from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Flowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table
from reportlab.platypus import TableStyle

from apps.finance.analytics.combined.exporters import (
    PDF_FONT_NAME,
    format_pdf_amount,
    register_pdf_font,
    truncate_text,
)
from apps.finance.analytics.forecast.services import (
    FORECAST_CHART_FOOTNOTE,
    FORECAST_DEFAULT_EXPORT_SECTIONS,
    FORECAST_DISCLAIMER,
    FORECAST_EXPORT_FORMATS,
    FORECAST_EXPORT_SECTIONS,
    METRIC_LABELS,
)


@dataclass(frozen=True)
class ForecastingExportResult:
    content: bytes
    content_type: str
    filename: str


CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

SECTION_LABELS = {
    "history": "История",
    "forecast": "Прогноз",
    "confidence": "Интервал доверия",
    "parameters": "Параметры модели",
}


class ForecastingPdfChart(Flowable):
    def __init__(self, *, projection: dict, font_name: str = PDF_FONT_NAME) -> None:
        super().__init__()
        self.width = 500
        self.height = 285
        self.projection = projection
        self.font_name = font_name

    def draw(self) -> None:
        canvas = self.canv
        self.draw_frame("История и прогноз")
        points = self.projection.get("chart_points", [])
        if not points:
            self.draw_text(24, self.height / 2, "Нет данных для отображения", size=10, color="#64748B")
            return

        chart_x = 58
        chart_y = 48
        chart_width = 340
        chart_height = 175
        legend_x = 418
        legend_y = 190

        values = []
        for point in points:
            values.append(float(point.get("value") or 0))
            if point.get("is_forecast"):
                values.append(float(point.get("lower_bound") or point.get("value") or 0))
                values.append(float(point.get("upper_bound") or point.get("value") or 0))
        min_value = min(min(values), 0)
        max_value = max(max(values), 0)
        if min_value == max_value:
            max_value = min_value + 1
        value_range = max_value - min_value

        def to_xy(index: int, value: float) -> tuple[float, float]:
            if len(points) == 1:
                x = chart_x + chart_width / 2
            else:
                x = chart_x + chart_width * index / (len(points) - 1)
            y = chart_y + chart_height * (value - min_value) / value_range
            return x, y

        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        canvas.setLineWidth(0.7)
        for tick in range(5):
            y = chart_y + chart_height * tick / 4
            tick_value = min_value + value_range * tick / 4
            canvas.line(chart_x, y, chart_x + chart_width, y)
            self.draw_right_text(chart_x - 7, y - 3, format_pdf_amount(tick_value), size=6, color="#64748B")

        forecast_indices = [index for index, point in enumerate(points) if point.get("is_forecast")]
        if forecast_indices:
            upper = [to_xy(index, float(points[index].get("upper_bound") or points[index].get("value") or 0)) for index in forecast_indices]
            lower = [to_xy(index, float(points[index].get("lower_bound") or points[index].get("value") or 0)) for index in reversed(forecast_indices)]
            polygon = upper + lower
            if len(polygon) >= 3:
                path = canvas.beginPath()
                path.moveTo(*polygon[0])
                for x, y in polygon[1:]:
                    path.lineTo(x, y)
                path.close()
                canvas.setFillColor(colors.HexColor("#DDD6FE"))
                canvas.drawPath(path, stroke=0, fill=1)
            elif len(forecast_indices) == 1:
                index = forecast_indices[0]
                lower_xy = to_xy(index, float(points[index].get("lower_bound") or points[index].get("value") or 0))
                upper_xy = to_xy(index, float(points[index].get("upper_bound") or points[index].get("value") or 0))
                canvas.setStrokeColor(colors.HexColor("#C4B5FD"))
                canvas.setLineWidth(5)
                canvas.line(lower_xy[0], lower_xy[1], upper_xy[0], upper_xy[1])

        history_coords = [
            to_xy(index, float(point.get("value") or 0))
            for index, point in enumerate(points)
            if not point.get("is_forecast")
        ]
        forecast_coords = []
        last_history_index = next((index for index in range(len(points) - 1, -1, -1) if not points[index].get("is_forecast")), None)
        if last_history_index is not None and forecast_indices:
            forecast_coords.append(to_xy(last_history_index, float(points[last_history_index].get("value") or 0)))
        forecast_coords.extend(
            to_xy(index, float(points[index].get("value") or 0))
            for index in forecast_indices
        )

        self.draw_line(history_coords, "#4F46E5", dashed=False)
        self.draw_line(forecast_coords, "#7C3AED", dashed=True)

        canvas.setStrokeColor(colors.HexColor("#94A3B8"))
        canvas.setLineWidth(1)
        canvas.setDash()
        canvas.line(chart_x, chart_y, chart_x + chart_width, chart_y)
        canvas.line(chart_x, chart_y, chart_x, chart_y + chart_height)

        label_step = max(len(points) // 6, 1)
        for index, point in enumerate(points):
            if index % label_step != 0 and index != len(points) - 1:
                continue
            x, _y = to_xy(index, min_value)
            self.draw_right_text(x + 17, chart_y - 15, truncate_text(point.get("label", ""), 9), size=6, color="#475569")

        self.draw_legend(legend_x, legend_y)

    def draw_line(self, coordinates: list[tuple[float, float]], color_hex: str, *, dashed: bool) -> None:
        if not coordinates:
            return
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor(color_hex))
        canvas.setFillColor(colors.HexColor(color_hex))
        canvas.setLineWidth(2)
        canvas.setDash(4, 3) if dashed else canvas.setDash()
        for start, end in zip(coordinates, coordinates[1:]):
            canvas.line(start[0], start[1], end[0], end[1])
        canvas.setDash()
        for x, y in coordinates:
            canvas.circle(x, y, 2.8, stroke=0, fill=1)

    def draw_frame(self, title: str) -> None:
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
        canvas.setLineWidth(0.8)
        canvas.roundRect(0, 0, self.width, self.height, 8, stroke=1, fill=0)
        self.draw_text(24, self.height - 28, title, size=11, color="#111827", bold=True)

    def draw_legend(self, x: float, y: float) -> None:
        items = [
            ("История", "#4F46E5"),
            ("Прогноз", "#7C3AED"),
            ("Интервал", "#DDD6FE"),
        ]
        canvas = self.canv
        for index, (label, color_hex) in enumerate(items):
            row_y = y - index * 19
            canvas.setFillColor(colors.HexColor(color_hex))
            canvas.rect(x, row_y, 9, 9, stroke=0, fill=1)
            self.draw_text(x + 14, row_y, label, size=8, color="#334155")

    def draw_text(self, x: float, y: float, text: object, *, size: int = 8, color: str = "#111827", bold: bool = False) -> None:
        canvas = self.canv
        canvas.setFillColor(colors.HexColor(color))
        canvas.setFont(self.font_name, size)
        canvas.drawString(x, y, str(text or ""))

    def draw_right_text(self, x: float, y: float, text: object, *, size: int = 8, color: str = "#111827") -> None:
        canvas = self.canv
        canvas.setFillColor(colors.HexColor(color))
        canvas.setFont(self.font_name, size)
        canvas.drawRightString(x, y, str(text or ""))


def build_forecasting_export(
    *,
    projection: dict,
    export_format: str,
    sections: list[str],
    metric: str,
) -> ForecastingExportResult:
    filename = build_forecasting_filename(export_format=export_format, metric=metric)
    if export_format == "csv":
        content = build_csv_export(projection=projection, sections=sections, metric=metric)
    elif export_format == "xlsx":
        content = build_xlsx_export(projection=projection, sections=sections, metric=metric)
    elif export_format == "pdf":
        content = build_pdf_export(projection=projection, sections=sections, metric=metric)
    else:
        raise ValueError(f"Unsupported export format: {export_format}")
    return ForecastingExportResult(
        content=content,
        content_type=CONTENT_TYPES[export_format],
        filename=filename,
    )


def build_forecasting_filename(*, export_format: str, metric: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_metric = str(metric or "forecast").lower().replace(" ", "_")
    return f"forecasting_analytics_{safe_metric}_{timestamp}.{export_format}"


def normalize_export_format(value: object | None) -> str:
    export_format = str(value or "").strip().lower()
    if export_format not in FORECAST_EXPORT_FORMATS:
        from rest_framework.exceptions import ValidationError

        raise ValidationError({"format": "Недопустимый формат экспорта."})
    return export_format


def normalize_export_sections(value: object | None) -> list[str]:
    if value is None:
        return list(FORECAST_DEFAULT_EXPORT_SECTIONS)
    if isinstance(value, str):
        raw_sections = [section.strip() for section in value.split(",") if section.strip()]
    elif isinstance(value, Iterable):
        raw_sections = [str(section).strip() for section in value if str(section).strip()]
    else:
        raw_sections = []
    if not raw_sections:
        from rest_framework.exceptions import ValidationError

        raise ValidationError({"sections": "Нужно выбрать хотя бы одну секцию экспорта."})
    invalid = [section for section in raw_sections if section not in FORECAST_EXPORT_SECTIONS]
    if invalid:
        from rest_framework.exceptions import ValidationError

        raise ValidationError({"sections": "Недопустимая секция экспорта."})
    result = []
    for section in raw_sections:
        if section not in result:
            result.append(section)
    return result


def build_csv_export(*, projection: dict, sections: list[str], metric: str) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(["Прогнозирование"])
    writer.writerow(["Метрика", METRIC_LABELS.get(metric, metric)])
    writer.writerow(["Есть данные", "Да" if projection.get("has_data") else "Нет"])
    writer.writerow(["Недостаточно данных", "Да" if projection.get("has_insufficient_data") else "Нет"])
    writer.writerow(["Высокая нестабильность", "Да" if projection.get("has_high_instability") else "Нет"])
    writer.writerow([])

    for section in sections:
        writer.writerow([SECTION_LABELS.get(section, section)])
        for row in build_section_rows(projection=projection, section=section):
            writer.writerow(row)
        writer.writerow([])
    return output.getvalue().encode("utf-8-sig")


def build_xlsx_export(*, projection: dict, sections: list[str], metric: str) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    write_summary_sheet(summary, projection=projection, metric=metric)

    if should_include_chart(sections):
        chart_sheet = workbook.create_sheet("График")
        write_chart_sheet(chart_sheet, projection=projection)
    if "history" in sections:
        history_sheet = workbook.create_sheet("История")
        write_section_sheet(history_sheet, projection=projection, section="history")
    if "forecast" in sections:
        forecast_sheet = workbook.create_sheet("Прогноз")
        write_section_sheet(forecast_sheet, projection=projection, section="forecast")
    if "confidence" in sections:
        confidence_sheet = workbook.create_sheet("Интервал")
        write_section_sheet(confidence_sheet, projection=projection, section="confidence")
    if "parameters" in sections:
        parameters_sheet = workbook.create_sheet("Параметры")
        write_section_sheet(parameters_sheet, projection=projection, section="parameters")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_summary_sheet(sheet, *, projection: dict, metric: str) -> None:
    rows = [
        ["Прогнозирование"],
        ["Метрика", METRIC_LABELS.get(metric, metric)],
        ["Есть данные", "Да" if projection.get("has_data") else "Нет"],
        ["Недостаточно данных", "Да" if projection.get("has_insufficient_data") else "Нет"],
        ["Высокая нестабильность", "Да" if projection.get("has_high_instability") else "Нет"],
        [],
        ["Суммарный прогноз", projection.get("metric_summary", {}).get("total_forecast_rub", 0)],
        ["Изменение, %", projection.get("metric_summary", {}).get("delta_percent", 0)],
        ["Пояснение", projection.get("metric_summary", {}).get("delta_label", "")],
    ]
    alert = projection.get("alert")
    if alert:
        rows.extend([
            [],
            ["Предупреждение", alert.get("title", "")],
            ["Описание", alert.get("text", "")],
        ])
    write_rows(sheet, rows)
    style_sheet(sheet)


def write_chart_sheet(sheet, *, projection: dict) -> None:
    rows = [["Месяц", "История", "Прогноз", "Нижняя граница", "Верхняя граница"]]
    for point in projection.get("chart_points", []):
        if point.get("is_forecast"):
            rows.append([
                point.get("label", ""),
                None,
                point.get("value", 0),
                point.get("lower_bound", 0),
                point.get("upper_bound", 0),
            ])
        else:
            rows.append([point.get("label", ""), point.get("value", 0), None, None, None])
    write_rows(sheet, rows)
    style_sheet(sheet)
    for row_number in range(2, sheet.max_row + 1):
        for column_number in range(2, 6):
            sheet.cell(row=row_number, column=column_number).number_format = '#,##0.00'
    if len(rows) <= 1:
        sheet.append(["Нет данных для построения графика"])
        return

    chart = LineChart()
    chart.title = "История и прогноз"
    chart.style = 13
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.y_axis.numFmt = '#,##0'
    chart.legend.position = "r"
    chart.legend.overlay = False
    data = Reference(sheet, min_col=2, max_col=5, min_row=1, max_row=sheet.max_row)
    cats = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 24
    chart.height = 13
    chart.layout = Layout(manualLayout=ManualLayout(x=0.06, y=0.08, w=0.72, h=0.80))
    sheet.add_chart(chart, "G2")


def write_section_sheet(sheet, *, projection: dict, section: str) -> None:
    rows = build_section_rows(projection=projection, section=section)
    write_rows(sheet, rows)
    if len(rows) == 1:
        sheet.append(["Нет данных"])
    style_sheet(sheet)


def build_section_rows(*, projection: dict, section: str) -> list[list[object]]:
    points = projection.get("chart_points", [])
    if section == "history":
        return [
            ["Месяц", "Значение"],
            *[
                [point.get("label", ""), point.get("value", 0)]
                for point in points
                if not point.get("is_forecast")
            ],
        ]
    if section == "forecast":
        return [
            ["Месяц", "Прогноз"],
            *[
                [row.get("month_label", ""), row.get("forecast_rub", 0)]
                for row in projection.get("detail_rows", [])
            ],
        ]
    if section == "confidence":
        return [
            ["Месяц", "Прогноз", "Нижняя граница", "Верхняя граница"],
            *[
                [row.get("month_label", ""), row.get("forecast_rub", 0), row.get("lower_rub", 0), row.get("upper_rub", 0)]
                for row in projection.get("detail_rows", [])
            ],
        ]
    if section == "parameters":
        assumptions = projection.get("assumptions", {})
        summary = projection.get("metric_summary", {})
        alert = projection.get("alert") or {}
        rows = [
            ["Параметр", "Значение"],
            ["Период истории", assumptions.get("history_period_label", "")],
            ["Источник", assumptions.get("source_label", "")],
            ["Модель", assumptions.get("model_label", "")],
            ["Последнее обновление", assumptions.get("updated_at_label", "")],
            ["Суммарный прогноз", summary.get("total_forecast_rub", 0)],
            ["Изменение, %", summary.get("delta_percent", 0)],
            ["Дисклеймер", FORECAST_DISCLAIMER],
            ["Сноска", FORECAST_CHART_FOOTNOTE],
        ]
        if alert:
            rows.append(["Предупреждение", f"{alert.get('title', '')}: {alert.get('text', '')}"])
        return rows
    return [["Нет данных"]]


def build_pdf_export(*, projection: dict, sections: list[str], metric: str) -> bytes:
    font_name = register_pdf_font()
    styles = getSampleStyleSheet()
    for style_name in ["Title", "Heading2", "BodyText"]:
        styles[style_name].fontName = font_name

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=32,
        leftMargin=32,
        topMargin=32,
        bottomMargin=32,
    )
    story = [Paragraph("Прогнозирование", styles["Title"]), Spacer(1, 8)]
    story.append(Paragraph(f"Метрика: {METRIC_LABELS.get(metric, metric)}", styles["BodyText"]))
    story.append(Paragraph(f"Есть данные: {'да' if projection.get('has_data') else 'нет'}", styles["BodyText"]))
    story.append(Paragraph(f"Недостаточно данных: {'да' if projection.get('has_insufficient_data') else 'нет'}", styles["BodyText"]))
    story.append(Paragraph(f"Высокая нестабильность: {'да' if projection.get('has_high_instability') else 'нет'}", styles["BodyText"]))
    story.append(Spacer(1, 12))

    alert = projection.get("alert")
    if alert:
        story.append(Paragraph(alert.get("title", "Предупреждение"), styles["Heading2"]))
        story.append(Paragraph(alert.get("text", ""), styles["BodyText"]))
        story.append(Spacer(1, 12))

    if should_include_chart(sections):
        story.append(ForecastingPdfChart(projection=projection, font_name=font_name))
        story.append(Spacer(1, 12))

    for section in sections:
        section_block = [
            Paragraph(SECTION_LABELS.get(section, section), styles["Heading2"]),
            build_forecasting_pdf_table(
                build_section_rows(projection=projection, section=section),
                section=section,
                font_name=font_name,
                available_width=doc.width,
            ),
            Spacer(1, 12),
        ]
        story.append(KeepTogether(section_block))

    doc.build(story)
    return buffer.getvalue()


def build_forecasting_pdf_table(
    rows: list[list[object]],
    *,
    section: str,
    font_name: str,
    available_width: float,
) -> Table:
    if len(rows) == 1:
        rows = [rows[0], ["Нет данных"]]

    header_style = ParagraphStyle(
        "ForecastingTableHeader",
        fontName=font_name,
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#111827"),
        wordWrap="CJK",
    )
    body_style = ParagraphStyle(
        "ForecastingTableBody",
        fontName=font_name,
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#111827"),
        wordWrap="CJK",
    )
    long_body_style = ParagraphStyle(
        "ForecastingTableBodyLong",
        fontName=font_name,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#111827"),
        wordWrap="CJK",
    )

    normalized_rows = normalize_pdf_table_rows(rows)
    prepared_rows = []
    for row_index, row in enumerate(normalized_rows):
        prepared_row = []
        for cell in row:
            style = header_style if row_index == 0 else body_style
            if row_index > 0 and isinstance(cell, str) and len(cell) > 80:
                style = long_body_style
            prepared_row.append(Paragraph(escape_pdf_text(cell), style))
        prepared_rows.append(prepared_row)

    table = Table(
        prepared_rows,
        colWidths=get_pdf_column_widths(
            section=section,
            column_count=len(prepared_rows[0]),
            available_width=available_width,
        ),
        repeatRows=1,
        hAlign="CENTER",
        splitByRow=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def normalize_pdf_table_rows(rows: list[list[object]]) -> list[list[object]]:
    max_columns = max((len(row) for row in rows), default=1)
    return [list(row) + [""] * (max_columns - len(row)) for row in rows]


def get_pdf_column_widths(*, section: str, column_count: int, available_width: float) -> list[float] | None:
    width = min(available_width, 500)
    if section == "parameters" and column_count == 2:
        return [130, width - 130]
    if section == "confidence" and column_count == 4:
        return [95, 120, 140, width - 355]
    if column_count == 2:
        return [130, 140]
    if column_count == 1:
        return [width]
    return [width / column_count] * column_count


def escape_pdf_text(value: object) -> str:
    text = str(value if value is not None else "")
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def should_include_chart(sections: list[str]) -> bool:
    return bool({"history", "forecast", "confidence"}.intersection(sections))


def write_rows(sheet, rows: list[list[object]]) -> None:
    for row in rows:
        sheet.append(row)


def style_sheet(sheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    title_fill = PatternFill(fill_type="solid", fgColor="EEF2FF")
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            if cell.row == 1:
                cell.font = Font(bold=True, color="111827")
                cell.fill = header_fill if sheet.max_column > 1 else title_fill
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        column_letter = column_cells[0].column_letter
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 42)
