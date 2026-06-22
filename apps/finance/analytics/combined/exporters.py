from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


@dataclass(frozen=True)
class CombinedAnalyticsExportResult:
    content: bytes
    content_type: str
    filename: str


PDF_FONT_NAME = "BudgetWiseDejaVu"
PDF_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
    "/usr/local/share/fonts/DejaVuSans.ttf",
]

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

SECTION_LABELS = {
    "pie": "Круговая диаграмма",
    "bar": "Столбчатая диаграмма",
    "line": "Линейный график",
    "table": "Таблица агрегатов",
}


def build_combined_analytics_export(
    *,
    aggregates: dict,
    export_format: str,
    sections: list[str],
) -> CombinedAnalyticsExportResult:
    filename = build_combined_analytics_filename(
        export_format=export_format,
        period_label=aggregates.get("period_label") or "analytics",
    )

    if export_format == "csv":
        content = build_csv_export(aggregates=aggregates, sections=sections)
    elif export_format == "xlsx":
        content = build_xlsx_export(aggregates=aggregates, sections=sections)
    elif export_format == "pdf":
        content = build_pdf_export(aggregates=aggregates, sections=sections)
    else:
        raise ValueError(f"Unsupported export format: {export_format}")

    return CombinedAnalyticsExportResult(
        content=content,
        content_type=CONTENT_TYPES[export_format],
        filename=filename,
    )


def build_combined_analytics_filename(*, export_format: str, period_label: str) -> str:
    safe_period = (
        str(period_label)
        .lower()
        .replace(" ", "_")
        .replace(".", "-")
        .replace("/", "-")
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"combined_analytics_{safe_period}_{timestamp}.{export_format}"


def build_csv_export(*, aggregates: dict, sections: list[str]) -> bytes:
    # Excel in Russian/European locales uses semicolon as the default CSV separator.
    # UTF-8 BOM keeps Cyrillic headers readable when the file is opened directly.
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")

    writer.writerow(["Комбинированная аналитика"])
    writer.writerow(["Период", aggregates.get("period_label", "")])
    writer.writerow(["Всего расходов", aggregates.get("total_expense_rub", 0)])
    writer.writerow(["Всего доходов", aggregates.get("total_income_rub", 0)])
    writer.writerow([])

    for section in sections:
        writer.writerow([SECTION_LABELS.get(section, section)])
        for row in build_section_rows(aggregates=aggregates, section=section):
            writer.writerow(row)
        writer.writerow([])

    return output.getvalue().encode("utf-8-sig")


def build_xlsx_export(*, aggregates: dict, sections: list[str]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    write_summary_sheet(summary_sheet, aggregates=aggregates)

    for section in sections:
        worksheet = workbook.create_sheet(title=SECTION_LABELS.get(section, section)[:31])
        write_section_sheet(worksheet, aggregates=aggregates, section=section)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_summary_sheet(worksheet, *, aggregates: dict) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A3"

    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "Комбинированная аналитика"
    worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor="1F2937")
    worksheet["A1"].alignment = Alignment(horizontal="center")

    rows = [
        ["Период", aggregates.get("period_label", "")],
        ["Всего расходов", aggregates.get("total_expense_rub", 0)],
        ["Всего доходов", aggregates.get("total_income_rub", 0)],
        ["Данные для отображения", "Да" if aggregates.get("has_data") else "Нет"],
    ]
    for row in rows:
        worksheet.append(row)

    style_range_as_report(worksheet, min_row=2, max_row=5, max_col=2)
    for cell in worksheet["A"]:
        if cell.row >= 2:
            cell.font = Font(bold=True)
    for row in range(3, 5):
        worksheet.cell(row=row, column=2).number_format = '#,##0.00'

    worksheet.column_dimensions["A"].width = 26
    worksheet.column_dimensions["B"].width = 22
    worksheet.column_dimensions["C"].width = 16
    worksheet.column_dimensions["D"].width = 16


def write_section_sheet(worksheet, *, aggregates: dict, section: str) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"

    rows = build_section_rows(aggregates=aggregates, section=section)
    has_data_rows = len(rows) > 1

    for row in rows:
        worksheet.append(row)

    if not has_data_rows:
        worksheet.append(["Нет данных за выбранный период"] + [None] * max(len(rows[0]) - 1, 0))

    max_row = worksheet.max_row
    max_col = worksheet.max_column
    style_header_row(worksheet, 1)
    style_range_as_report(worksheet, min_row=1, max_row=max_row, max_col=max_col)
    apply_section_number_formats(worksheet, section=section)
    autosize_columns(worksheet)

    if has_data_rows:
        worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=max_row, column=max_col).coordinate}"
        add_section_chart(worksheet, section=section)



def add_section_chart(worksheet, *, section: str) -> None:
    if worksheet.max_row < 2:
        return

    # XLSX is used as a visual report. Excel can easily make labels overlap
    # when charts are created with default settings, especially with long
    # Russian category names. Therefore charts below intentionally keep the
    # plot area clean: no axis titles, no data labels on crowded charts,
    # a non-overlapping legend, larger chart size, and a fixed inner layout.
    if section == "pie":
        chart = PieChart()
        labels = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row)
        data = Reference(worksheet, min_col=2, min_row=1, max_row=worksheet.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        chart.title = "Расходы по категориям"
        chart.style = 10
        chart.height = 14
        chart.width = 24
        chart.legend.position = "r"
        chart.legend.overlay = False
        chart.firstSliceAng = 270
        chart.layout = Layout(
            manualLayout=ManualLayout(
                x=0.03,
                y=0.08,
                w=0.68,
                h=0.82,
            )
        )
        worksheet.add_chart(chart, "E2")
        return

    if section == "bar":
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Сравнение периодов"
        chart.legend.position = "r"
        chart.legend.overlay = False
        chart.y_axis.numFmt = '#,##0'
        chart.x_axis.tickLblPos = "low"
        chart.y_axis.tickLblPos = "low"
        data = Reference(
            worksheet,
            min_col=2,
            max_col=3,
            min_row=1,
            max_row=worksheet.max_row,
        )
        categories = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 14
        chart.width = 26
        chart.layout = Layout(
            manualLayout=ManualLayout(
                x=0.10,
                y=0.10,
                w=0.68,
                h=0.78,
            )
        )
        worksheet.add_chart(chart, "E2")
        return

    if section == "line":
        chart = LineChart()
        chart.style = 13
        chart.title = "Финансовые тренды"
        chart.legend.position = "r"
        chart.legend.overlay = False
        chart.y_axis.numFmt = '#,##0'
        chart.x_axis.tickLblPos = "low"
        data = Reference(
            worksheet,
            min_col=2,
            max_col=4,
            min_row=1,
            max_row=worksheet.max_row,
        )
        categories = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 14
        chart.width = 26
        chart.layout = Layout(
            manualLayout=ManualLayout(
                x=0.08,
                y=0.10,
                w=0.70,
                h=0.78,
            )
        )
        worksheet.add_chart(chart, "F2")


def apply_section_number_formats(worksheet, *, section: str) -> None:
    if worksheet.max_row < 2:
        return

    if section in {"pie", "table"}:
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=2).number_format = '#,##0.00'
            worksheet.cell(row=row, column=3).number_format = '0.00'

    if section == "bar":
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=2).number_format = '#,##0.00'
            worksheet.cell(row=row, column=3).number_format = '#,##0.00'

    if section == "line":
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=2).number_format = '#,##0.00'
            worksheet.cell(row=row, column=3).number_format = '#,##0.00'
            worksheet.cell(row=row, column=4).number_format = '#,##0.00'


def build_pdf_export(*, aggregates: dict, sections: list[str]) -> bytes:
    font_name = register_pdf_font()

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=28,
    )
    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]
    title_style = styles["Title"]
    heading_style = styles["Heading2"]

    for style in (normal_style, title_style, heading_style):
        style.fontName = font_name

    story = [
        Paragraph("Комбинированная аналитика", title_style),
        Paragraph(f"Период: {aggregates.get('period_label', '')}", normal_style),
        Paragraph(f"Всего расходов: {format_pdf_amount(aggregates.get('total_expense_rub', 0))}", normal_style),
        Paragraph(f"Всего доходов: {format_pdf_amount(aggregates.get('total_income_rub', 0))}", normal_style),
        Spacer(1, 12),
    ]

    for section in sections:
        rows = build_section_rows(aggregates=aggregates, section=section)
        story.append(Paragraph(SECTION_LABELS.get(section, section), heading_style))

        chart = build_pdf_chart(aggregates=aggregates, section=section, font_name=font_name)
        if chart is not None:
            story.append(chart)
            story.append(Spacer(1, 10))

        story.append(build_pdf_table(rows, font_name=font_name))
        story.append(Spacer(1, 14))

    document.build(story)
    return buffer.getvalue()


class CombinedAnalyticsChartFlowable(Flowable):
    """Simple ReportLab chart flowable for the PDF export.

    The XLSX export uses native Excel charts, but PDF needs charts rendered into
    the document itself. These small canvas-based charts keep labels outside the
    plot area, so Russian category names do not overlap bars and lines.
    """

    def __init__(self, *, width: float = 500, height: float = 245, font_name: str = PDF_FONT_NAME) -> None:
        super().__init__()
        self.width = width
        self.height = height
        self.font_name = font_name

    def wrap(self, availWidth, availHeight):
        return min(self.width, availWidth), self.height

    def draw_frame(self, title: str) -> None:
        canvas = self.canv
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
        canvas.setFillColor(colors.HexColor("#F8FAFC"))
        canvas.roundRect(0, 0, self.width, self.height, 8, stroke=1, fill=1)
        canvas.setFillColor(colors.HexColor("#111827"))
        canvas.setFont(self.font_name, 12)
        canvas.drawCentredString(self.width / 2, self.height - 22, title)
        canvas.restoreState()

    def draw_text(self, x: float, y: float, text: object, *, size: int = 8, color="#111827") -> None:
        canvas = self.canv
        canvas.setFillColor(colors.HexColor(color))
        canvas.setFont(self.font_name, size)
        canvas.drawString(x, y, str(text))

    def draw_right_text(self, x: float, y: float, text: object, *, size: int = 8, color="#111827") -> None:
        canvas = self.canv
        canvas.setFillColor(colors.HexColor(color))
        canvas.setFont(self.font_name, size)
        canvas.drawRightString(x, y, str(text))


class PiePdfChart(CombinedAnalyticsChartFlowable):
    def __init__(self, *, slices: list[dict], font_name: str = PDF_FONT_NAME) -> None:
        super().__init__(width=500, height=245, font_name=font_name)
        self.slices = slices

    def draw(self) -> None:
        self.draw_frame("Расходы по категориям")
        canvas = self.canv
        values = [float(item.get("amount_rub") or 0) for item in self.slices]
        total = sum(values)
        if not self.slices or total <= 0:
            self.draw_text(26, self.height / 2, "Нет данных за выбранный период", size=10, color="#64748B")
            return

        palette = get_chart_palette()
        cx, cy, radius = 135, 120, 70
        start_angle = 90
        for index, value in enumerate(values):
            extent = 360 * value / total
            canvas.setFillColor(colors.HexColor(palette[index % len(palette)]))
            canvas.setStrokeColor(colors.white)
            canvas.wedge(
                cx - radius,
                cy - radius,
                cx + radius,
                cy + radius,
                start_angle,
                extent,
                stroke=1,
                fill=1,
            )
            start_angle += extent

        legend_x = 250
        legend_y = 178
        row_height = 22
        self.draw_text(legend_x, legend_y + 22, "Категория", size=9, color="#334155")
        self.draw_right_text(self.width - 28, legend_y + 22, "Сумма / доля", size=9, color="#334155")
        for index, item in enumerate(self.slices[:8]):
            y = legend_y - index * row_height
            canvas.setFillColor(colors.HexColor(palette[index % len(palette)]))
            canvas.rect(legend_x, y - 1, 9, 9, stroke=0, fill=1)
            label = truncate_text(item.get("category_name", ""), 26)
            amount = format_pdf_amount(item.get("amount_rub", 0))
            percent = format_percent(item.get("percent", 0))
            self.draw_text(legend_x + 14, y, label, size=8)
            self.draw_right_text(self.width - 28, y, f"{amount} / {percent}%", size=8)


class BarPdfChart(CombinedAnalyticsChartFlowable):
    def __init__(self, *, groups: list[dict], legend: dict, font_name: str = PDF_FONT_NAME) -> None:
        super().__init__(width=500, height=270, font_name=font_name)
        self.groups = groups
        self.legend = legend

    def draw(self) -> None:
        self.draw_frame("Сравнение периодов")
        canvas = self.canv
        groups = self.groups[:8]
        if not groups:
            self.draw_text(26, self.height / 2, "Нет данных за выбранный период", size=10, color="#64748B")
            return

        previous_label = self.legend.get("previous_period_label") or "Предыдущий период"
        current_label = self.legend.get("current_period_label") or "Текущий период"
        max_value = max(
            max(float(item.get("previous_period_amount_rub") or 0), float(item.get("current_period_amount_rub") or 0))
            for item in groups
        )
        if max_value <= 0:
            max_value = 1

        palette = get_chart_palette()
        previous_color = colors.HexColor(palette[0])
        current_color = colors.HexColor(palette[1])
        label_x = 22
        chart_x = 180
        chart_y = 42
        chart_width = 260
        chart_height = 170
        group_gap = chart_height / max(len(groups), 1)
        bar_height = min(10, max(6, group_gap / 3))

        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        for tick in range(5):
            x = chart_x + chart_width * tick / 4
            canvas.line(x, chart_y, x, chart_y + chart_height)
            tick_value = max_value * tick / 4
            self.draw_right_text(x + 14, chart_y - 14, format_pdf_amount(tick_value), size=6, color="#64748B")

        self.draw_legend(previous_label, current_label, previous_color, current_color)

        for index, item in enumerate(groups):
            base_y = chart_y + chart_height - (index + 0.65) * group_gap
            label = truncate_text(item.get("category_name", ""), 28)
            self.draw_text(label_x, base_y + 1, label, size=8)

            previous_value = float(item.get("previous_period_amount_rub") or 0)
            current_value = float(item.get("current_period_amount_rub") or 0)
            previous_width = chart_width * previous_value / max_value
            current_width = chart_width * current_value / max_value

            canvas.setFillColor(previous_color)
            canvas.rect(chart_x, base_y + bar_height / 2, previous_width, bar_height, stroke=0, fill=1)
            canvas.setFillColor(current_color)
            canvas.rect(chart_x, base_y - bar_height, current_width, bar_height, stroke=0, fill=1)

        canvas.setStrokeColor(colors.HexColor("#94A3B8"))
        canvas.line(chart_x, chart_y, chart_x + chart_width, chart_y)
        canvas.line(chart_x, chart_y, chart_x, chart_y + chart_height)

    def draw_legend(self, previous_label, current_label, previous_color, current_color) -> None:
        canvas = self.canv
        legend_y = self.height - 50
        legend_x = 285
        canvas.setFillColor(previous_color)
        canvas.rect(legend_x, legend_y, 9, 9, stroke=0, fill=1)
        self.draw_text(legend_x + 14, legend_y, truncate_text(previous_label, 18), size=8)
        canvas.setFillColor(current_color)
        canvas.rect(legend_x + 110, legend_y, 9, 9, stroke=0, fill=1)
        self.draw_text(legend_x + 124, legend_y, truncate_text(current_label, 18), size=8)


class LinePdfChart(CombinedAnalyticsChartFlowable):
    def __init__(self, *, points: list[dict], font_name: str = PDF_FONT_NAME) -> None:
        super().__init__(width=500, height=270, font_name=font_name)
        self.points = points

    def draw(self) -> None:
        self.draw_frame("Финансовые тренды")
        points = self.points[:12]
        if not points:
            self.draw_text(26, self.height / 2, "Нет данных за выбранный период", size=10, color="#64748B")
            return

        series = [
            ("Доходы", [float(item.get("income_rub") or 0) for item in points], "#22C55E"),
            ("Расходы", [float(item.get("expense_rub") or 0) for item in points], "#EF4444"),
            ("Баланс", [float(item.get("balance_rub") or 0) for item in points], "#3B82F6"),
        ]
        all_values = [value for _, values, _ in series for value in values]
        min_value = min(min(all_values), 0)
        max_value = max(max(all_values), 0)
        if min_value == max_value:
            max_value = min_value + 1

        canvas = self.canv
        chart_x = 58
        chart_y = 52
        chart_width = 350
        chart_height = 165
        value_range = max_value - min_value

        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        for tick in range(5):
            y = chart_y + chart_height * tick / 4
            value = min_value + value_range * tick / 4
            canvas.line(chart_x, y, chart_x + chart_width, y)
            self.draw_right_text(chart_x - 7, y - 3, format_pdf_amount(value), size=6, color="#64748B")

        def point_coordinates(values: list[float]) -> list[tuple[float, float]]:
            if len(values) == 1:
                return [(chart_x + chart_width / 2, chart_y + chart_height * (values[0] - min_value) / value_range)]
            return [
                (
                    chart_x + chart_width * index / (len(values) - 1),
                    chart_y + chart_height * (value - min_value) / value_range,
                )
                for index, value in enumerate(values)
            ]

        for label, values, color_hex in series:
            coordinates = point_coordinates(values)
            canvas.setStrokeColor(colors.HexColor(color_hex))
            canvas.setLineWidth(2)
            for start, end in zip(coordinates, coordinates[1:]):
                canvas.line(start[0], start[1], end[0], end[1])
            canvas.setFillColor(colors.HexColor(color_hex))
            for x, y in coordinates:
                canvas.circle(x, y, 2.6, stroke=0, fill=1)

        canvas.setStrokeColor(colors.HexColor("#94A3B8"))
        canvas.setLineWidth(1)
        canvas.line(chart_x, chart_y, chart_x + chart_width, chart_y)
        canvas.line(chart_x, chart_y, chart_x, chart_y + chart_height)

        labels = [truncate_text(item.get("label", ""), 9) for item in points]
        for index, label in enumerate(labels):
            x = chart_x + (chart_width * index / max(len(labels) - 1, 1) if len(labels) > 1 else chart_width / 2)
            self.draw_right_text(x + 16, chart_y - 16, label, size=6, color="#475569")

        self.draw_legend(series)

    def draw_legend(self, series: list[tuple[str, list[float], str]]) -> None:
        canvas = self.canv
        legend_x = 425
        legend_y = 170
        for index, (label, _values, color_hex) in enumerate(series):
            y = legend_y - index * 18
            canvas.setFillColor(colors.HexColor(color_hex))
            canvas.rect(legend_x, y, 9, 9, stroke=0, fill=1)
            self.draw_text(legend_x + 14, y, label, size=8)


def build_pdf_chart(*, aggregates: dict, section: str, font_name: str) -> Flowable | None:
    if section == "pie":
        return PiePdfChart(slices=aggregates.get("pie_slices", []), font_name=font_name)
    if section == "bar":
        return BarPdfChart(
            groups=aggregates.get("bar_groups", []),
            legend=aggregates.get("bar_legend", {}),
            font_name=font_name,
        )
    if section == "line":
        return LinePdfChart(points=aggregates.get("line_points", []), font_name=font_name)
    return None


def get_chart_palette() -> list[str]:
    return [
        "#4F46E5",
        "#EF4444",
        "#22C55E",
        "#F97316",
        "#8B5CF6",
        "#06B6D4",
        "#EC4899",
        "#84CC16",
    ]


def format_pdf_amount(value: object) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        return "0"
    formatted = f"{number:,.0f}" if number == int(number) else f"{number:,.2f}"
    return formatted.replace(",", " ")


def format_percent(value: object) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        return "0"
    return f"{number:.2f}".rstrip("0").rstrip(".")


def truncate_text(value: object, max_length: int) -> str:
    text = str(value or "")
    if len(text) <= max_length:
        return text
    return text[: max(max_length - 3, 1)] + "..."


def build_section_rows(*, aggregates: dict, section: str) -> list[list[object]]:
    if section == "pie":
        return [
            ["Категория", "Сумма", "Доля, %"],
            *[
                [row["category_name"], row["amount_rub"], row["percent"]]
                for row in aggregates.get("pie_slices", [])
            ],
        ]

    if section == "bar":
        legend = aggregates.get("bar_legend", {})
        previous_label = legend.get("previous_period_label", "Предыдущий период")
        current_label = legend.get("current_period_label", "Текущий период")
        return [
            ["Категория", previous_label, current_label],
            *[
                [
                    row["category_name"],
                    row["previous_period_amount_rub"],
                    row["current_period_amount_rub"],
                ]
                for row in aggregates.get("bar_groups", [])
            ],
        ]

    if section == "line":
        return [
            ["Месяц", "Доходы", "Расходы", "Баланс"],
            *[
                [row["label"], row["income_rub"], row["expense_rub"], row["balance_rub"]]
                for row in aggregates.get("line_points", [])
            ],
        ]

    if section == "table":
        return [
            ["Категория", "Сумма", "Доля, %"],
            *[
                [row["category_name"], row["amount_rub"], row["percent"]]
                for row in aggregates.get("aggregate_rows", [])
            ],
        ]

    return [["Нет данных"]]


def style_header_row(worksheet, row_number: int) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="111827")
    for cell in worksheet[row_number]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_range_as_report(worksheet, *, min_row: int, max_row: int, max_col: int) -> None:
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for row_idx in range(min_row, max_row + 1):
        worksheet.row_dimensions[row_idx].height = 22


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 36)


def build_pdf_table(rows: list[list[object]], *, font_name: str = PDF_FONT_NAME) -> Table:
    if len(rows) == 1:
        rows = [rows[0], ["Нет данных"]]

    table = Table(rows, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def register_pdf_font() -> str:
    registered_fonts = set(pdfmetrics.getRegisteredFontNames())
    if PDF_FONT_NAME in registered_fonts:
        return PDF_FONT_NAME

    for font_path in PDF_FONT_CANDIDATES:
        if Path(font_path).exists():
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, font_path))
            return PDF_FONT_NAME

    return "Helvetica"
