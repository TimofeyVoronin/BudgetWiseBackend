from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


@dataclass(frozen=True)
class ComparativeAnalyticsExportResult:
    content: bytes
    content_type: str
    filename: str


CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

SECTION_LABELS = {
    "legend": "Легенда сравнения",
    "differences": "Различия",
    "charts": "Графики",
}

PDF_FONT_NAME = "BudgetWiseDejaVu"
PDF_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
    "/usr/local/share/fonts/DejaVuSans.ttf",
]


def build_comparative_analytics_export(
    *,
    comparison: dict,
    export_format: str,
    sections: list[str],
) -> ComparativeAnalyticsExportResult:
    filename = build_filename(export_format=export_format)
    if export_format == "csv":
        content = build_csv_export(comparison=comparison, sections=sections)
    elif export_format == "xlsx":
        content = build_xlsx_export(comparison=comparison, sections=sections)
    elif export_format == "pdf":
        content = build_pdf_export(comparison=comparison, sections=sections)
    else:
        raise ValueError(f"Unsupported export format: {export_format}")
    return ComparativeAnalyticsExportResult(
        content=content,
        content_type=CONTENT_TYPES[export_format],
        filename=filename,
    )


def build_filename(*, export_format: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"comparative_analytics_{timestamp}.{export_format}"


def build_csv_export(*, comparison: dict, sections: list[str]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(["Сравнительная аналитика"])
    writer.writerow(["Есть данные", "Да" if comparison.get("has_data") else "Нет"])
    writer.writerow(["Разные валюты", "Да" if comparison.get("has_mixed_currencies") else "Нет"])
    writer.writerow([])

    if "legend" in sections:
        writer.writerow([SECTION_LABELS["legend"]])
        writer.writerow(["ID", "Название", "Описание", "Сумма", "Валюта", "Есть данные"])
        for entity in comparison.get("entities", []):
            writer.writerow([
                entity.get("id", ""),
                entity.get("label", ""),
                entity.get("subtitle", ""),
                entity.get("amount_rub", 0),
                entity.get("currency_code", ""),
                "Да" if entity.get("has_data") else "Нет",
            ])
        writer.writerow([])

    if "charts" in sections:
        writer.writerow(["Карточки сводки"])
        writer.writerow(["Показатель", "Сумма", "Изменение, %", "Пояснение"])
        for card in comparison.get("summary_cards", []):
            writer.writerow([
                card.get("label", ""),
                card.get("amount_rub", 0),
                card.get("delta_percent", 0),
                card.get("delta_label", ""),
            ])
        writer.writerow([])

        writer.writerow(["Динамика показателя"])
        entity_labels = [entity.get("label", "") for entity in comparison.get("entities", [])]
        writer.writerow(["Показатель", *entity_labels])
        for group in comparison.get("bar_groups", []):
            writer.writerow([group.get("label", ""), *group.get("values", [])])
        writer.writerow([])

        writer.writerow(["Динамика доходов и расходов"])
        writer.writerow(["Месяц", "Доходы", "Расходы"])
        for point in comparison.get("line_points", []):
            writer.writerow([point.get("label", ""), point.get("income_rub", 0), point.get("expense_rub", 0)])
        writer.writerow([])

    if "differences" in sections:
        writer.writerow([SECTION_LABELS["differences"]])
        writer.writerow(["Показатель", "Значение A", "Значение B", "Разница", "Разница, %"])
        for row in comparison.get("difference_rows", []):
            writer.writerow([
                row.get("metric", ""),
                row.get("value_a", 0),
                row.get("value_b", 0),
                row.get("delta_rub", 0),
                row.get("delta_percent", 0),
            ])
        writer.writerow([])

        writer.writerow(["Сравнение по категориям"])
        writer.writerow(["Категория", "Предыдущий период", "Текущий период", "Изменение, %"])
        for row in comparison.get("category_rows", []):
            writer.writerow([
                row.get("name", ""),
                row.get("previous_amount_rub", 0),
                row.get("current_amount_rub", 0),
                row.get("delta_percent", 0),
            ])

    return output.getvalue().encode("utf-8-sig")


def build_xlsx_export(*, comparison: dict, sections: list[str]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    write_summary_sheet(summary, comparison=comparison)

    if "legend" in sections:
        sheet = workbook.create_sheet("Легенда")
        write_entities_sheet(sheet, comparison=comparison)
    if "charts" in sections:
        sheet = workbook.create_sheet("Графики")
        write_charts_sheet(sheet, comparison=comparison)
    if "differences" in sections:
        sheet = workbook.create_sheet("Различия")
        write_differences_sheet(sheet, comparison=comparison)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_summary_sheet(sheet, *, comparison: dict) -> None:
    rows = [
        ["Сравнительная аналитика"],
        ["Есть данные", "Да" if comparison.get("has_data") else "Нет"],
        ["Разные валюты", "Да" if comparison.get("has_mixed_currencies") else "Нет"],
        [],
        ["Показатель", "Сумма", "Изменение, %", "Пояснение"],
    ]
    for card in comparison.get("summary_cards", []):
        rows.append([card.get("label"), card.get("amount_rub"), card.get("delta_percent"), card.get("delta_label")])
    write_rows(sheet, rows)
    style_sheet(sheet)


def write_entities_sheet(sheet, *, comparison: dict) -> None:
    rows = [["ID", "Название", "Описание", "Сумма", "Валюта", "Есть данные"]]
    for entity in comparison.get("entities", []):
        rows.append([
            entity.get("id"),
            entity.get("label"),
            entity.get("subtitle"),
            entity.get("amount_rub"),
            entity.get("currency_code"),
            "Да" if entity.get("has_data") else "Нет",
        ])
    write_rows(sheet, rows)
    style_sheet(sheet)


def write_charts_sheet(sheet, *, comparison: dict) -> None:
    entity_labels = [entity.get("label", "") for entity in comparison.get("entities", [])]
    bar_start = 1
    bar_rows = [["Показатель", *entity_labels]]
    for group in comparison.get("bar_groups", []):
        bar_rows.append([group.get("label", ""), *group.get("values", [])])
    write_rows(sheet, bar_rows, start_row=bar_start)

    # Keep enough vertical space between Excel chart objects.  With the previous
    # layout the second chart was anchored too close to the first one and Excel
    # rendered them on top of each other on some systems.
    line_start = max(bar_start + len(bar_rows) + 4, 28)
    line_rows = [["Месяц", "Доходы", "Расходы"]]
    for point in get_visible_line_points(comparison):
        line_rows.append([point.get("label", ""), point.get("income_rub", 0), point.get("expense_rub", 0)])
    write_rows(sheet, line_rows, start_row=line_start)
    style_sheet(sheet)

    if len(bar_rows) > 1 and len(entity_labels) > 0:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Динамика показателя"
        chart.y_axis.title = None
        chart.x_axis.title = None
        chart.width = 20
        chart.height = 10
        data = Reference(sheet, min_col=2, max_col=1 + len(entity_labels), min_row=bar_start, max_row=bar_start + len(bar_rows) - 1)
        cats = Reference(sheet, min_col=1, min_row=bar_start + 1, max_row=bar_start + len(bar_rows) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend.position = "r"
        chart.layout = Layout(manualLayout=ManualLayout(x=0.04, y=0.08, w=0.76, h=0.78))
        sheet.add_chart(chart, "H1")

    if len(line_rows) > 1:
        chart = LineChart()
        chart.title = "Динамика доходов и расходов"
        chart.y_axis.title = None
        chart.x_axis.title = None
        chart.width = 20
        chart.height = 10
        data = Reference(sheet, min_col=2, max_col=3, min_row=line_start, max_row=line_start + len(line_rows) - 1)
        cats = Reference(sheet, min_col=1, min_row=line_start + 1, max_row=line_start + len(line_rows) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend.position = "r"
        chart.layout = Layout(manualLayout=ManualLayout(x=0.04, y=0.08, w=0.76, h=0.78))
        sheet.add_chart(chart, f"H{line_start}")



def get_visible_line_points(comparison: dict) -> list[dict]:
    """Return line-chart points without trailing empty months.

    The API can still return a full calendar year.  For export charts we trim
    trailing months where both income and expense are zero, otherwise Excel
    draws a sharp artificial drop to zero after the last real data month.
    """
    points = list(comparison.get("line_points", []))
    last_index = -1
    for index, point in enumerate(points):
        income = float(point.get("income_rub") or 0)
        expense = float(point.get("expense_rub") or 0)
        if income != 0 or expense != 0:
            last_index = index
    return points if last_index < 0 else points[: last_index + 1]

def write_differences_sheet(sheet, *, comparison: dict) -> None:
    rows = [["Различия"]]
    rows.append(["Показатель", "Значение A", "Значение B", "Разница", "Разница, %", "A", "B"])
    for row in comparison.get("difference_rows", []):
        rows.append([
            row.get("metric"),
            row.get("value_a"),
            row.get("value_b"),
            row.get("delta_rub"),
            row.get("delta_percent"),
            row.get("label_a", ""),
            row.get("label_b", ""),
        ])
    rows.append([])
    rows.append(["Сравнение по категориям"])
    rows.append(["Категория", "Предыдущий период", "Текущий период", "Изменение, %"])
    for row in comparison.get("category_rows", []):
        rows.append([
            row.get("name"),
            row.get("previous_amount_rub"),
            row.get("current_amount_rub"),
            row.get("delta_percent"),
        ])
    write_rows(sheet, rows)
    style_sheet(sheet)


def write_rows(sheet, rows: list[list], *, start_row: int = 1, start_col: int = 1) -> None:
    for row_index, row in enumerate(rows, start=start_row):
        for col_index, value in enumerate(row, start=start_col):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            if isinstance(value, (int, float)):
                cell.number_format = '# ##0.00'


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    title_fill = PatternFill("solid", fgColor="EEF2FF")
    border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.fill = title_fill
            elif cell.value in {"Показатель", "ID", "Месяц", "Категория"} or cell.row in {5}:
                cell.font = Font(bold=True)
                cell.fill = header_fill
    widths = {"A": 28, "B": 18, "C": 18, "D": 18, "E": 16, "F": 24, "G": 24}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"


def build_pdf_export(*, comparison: dict, sections: list[str]) -> bytes:
    font_name = register_pdf_font()
    styles = getSampleStyleSheet()
    for style_name in ["Title", "Heading2", "BodyText"]:
        styles[style_name].fontName = font_name
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=32,
        leftMargin=32,
        topMargin=32,
        bottomMargin=32,
    )
    story = [Paragraph("Сравнительная аналитика", styles["Title"]), Spacer(1, 10)]
    story.append(Paragraph(f"Есть данные: {'да' if comparison.get('has_data') else 'нет'}", styles["BodyText"]))
    story.append(Paragraph(f"Разные валюты: {'да' if comparison.get('has_mixed_currencies') else 'нет'}", styles["BodyText"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Сводка", styles["Heading2"]))
    story.append(build_pdf_table(
        [["Показатель", "Сумма", "Изменение", "Пояснение"]]
        + [[card.get("label"), card.get("amount_rub"), card.get("delta_percent"), card.get("delta_label")] for card in comparison.get("summary_cards", [])],
        font_name=font_name,
    ))
    story.append(Spacer(1, 12))

    if "legend" in sections:
        story.append(Paragraph("Легенда сравнения", styles["Heading2"]))
        story.append(build_pdf_table(
            [["Название", "Описание", "Сумма", "Валюта"]]
            + [[entity.get("label"), entity.get("subtitle"), entity.get("amount_rub"), entity.get("currency_code")] for entity in comparison.get("entities", [])],
            font_name=font_name,
        ))
        story.append(Spacer(1, 12))

    if "charts" in sections:
        story.append(Paragraph("Графики", styles["Heading2"]))
        story.append(ComparativeBarPdfChart(comparison=comparison, font_name=font_name))
        story.append(Spacer(1, 8))
        story.append(ComparativeLinePdfChart(comparison=comparison, font_name=font_name))
        story.append(Spacer(1, 12))

    if "differences" in sections:
        story.append(Paragraph("Различия", styles["Heading2"]))
        story.append(build_pdf_table(
            [["Показатель", "A", "B", "Разница", "Разница, %"]]
            + [[row.get("metric"), row.get("value_a"), row.get("value_b"), row.get("delta_rub"), row.get("delta_percent")] for row in comparison.get("difference_rows", [])],
            font_name=font_name,
        ))
        story.append(Spacer(1, 10))
        story.append(Paragraph("Сравнение по категориям", styles["Heading2"]))
        story.append(build_pdf_table(
            [["Категория", "Предыдущий", "Текущий", "Изменение, %"]]
            + [[row.get("name"), row.get("previous_amount_rub"), row.get("current_amount_rub"), row.get("delta_percent")] for row in comparison.get("category_rows", [])],
            font_name=font_name,
        ))

    doc.build(story)
    return buffer.getvalue()


def build_pdf_table(rows: list[list], *, font_name: str) -> Table:
    if len(rows) <= 1:
        rows = rows + [["Нет данных", "", "", ""]]
    table = Table(rows, hAlign="CENTER")
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


class ComparativeChartFlowable(Flowable):
    def __init__(self, *, width=500, height=250, font_name=PDF_FONT_NAME) -> None:
        super().__init__()
        self.width = width
        self.height = height
        self.font_name = font_name

    def draw_frame(self, title: str) -> None:
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
        canvas.roundRect(0, 0, self.width, self.height, 8, stroke=1, fill=0)
        canvas.setFont(self.font_name, 10)
        canvas.setFillColor(colors.HexColor("#111827"))
        canvas.drawCentredString(self.width / 2, self.height - 20, title)

    def draw_text(self, x, y, text, *, size=8, color="#111827") -> None:
        self.canv.setFont(self.font_name, size)
        self.canv.setFillColor(colors.HexColor(color))
        self.canv.drawString(x, y, str(text))

    def draw_right_text(self, x, y, text, *, size=8, color="#111827") -> None:
        self.canv.setFont(self.font_name, size)
        self.canv.setFillColor(colors.HexColor(color))
        self.canv.drawRightString(x, y, str(text))


class ComparativeBarPdfChart(ComparativeChartFlowable):
    def __init__(self, *, comparison: dict, font_name: str = PDF_FONT_NAME) -> None:
        super().__init__(width=500, height=245, font_name=font_name)
        self.comparison = comparison

    def draw(self) -> None:
        self.draw_frame("Динамика показателя")
        groups = self.comparison.get("bar_groups", [])
        entities = self.comparison.get("entities", [])
        if not groups or not entities:
            self.draw_text(24, 115, "Нет данных за выбранный период", color="#64748B")
            return
        palette = [entity.get("color", "#4F46E5") for entity in entities]
        max_value = max([abs(float(value or 0)) for group in groups for value in group.get("values", [])] + [1])
        chart_x = 70
        chart_y = 46
        chart_width = 320
        chart_height = 150
        group_gap = chart_width / max(len(groups), 1)
        bar_width = min(18, group_gap / max(len(entities) + 1, 1))
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        for tick in range(5):
            y = chart_y + chart_height * tick / 4
            canvas.line(chart_x, y, chart_x + chart_width, y)
            self.draw_right_text(chart_x - 6, y - 3, format_amount(max_value * tick / 4), size=6, color="#64748B")
        for group_index, group in enumerate(groups):
            base_x = chart_x + group_gap * group_index + group_gap / 2
            values = group.get("values", [])
            for entity_index, value in enumerate(values):
                height = chart_height * abs(float(value or 0)) / max_value
                x = base_x + (entity_index - len(values) / 2) * bar_width
                canvas.setFillColor(colors.HexColor(palette[entity_index % len(palette)]))
                canvas.rect(x, chart_y, bar_width * 0.82, height, stroke=0, fill=1)
            self.draw_right_text(base_x + 12, chart_y - 14, group.get("label", ""), size=7, color="#475569")
        self.draw_legend(entities)

    def draw_legend(self, entities: list[dict]) -> None:
        y = 166
        x = 405
        for index, entity in enumerate(entities[:5]):
            row_y = y - index * 18
            self.canv.setFillColor(colors.HexColor(entity.get("color", "#4F46E5")))
            self.canv.rect(x, row_y, 8, 8, stroke=0, fill=1)
            self.draw_text(x + 12, row_y, truncate(entity.get("label", ""), 18), size=7)


class ComparativeLinePdfChart(ComparativeChartFlowable):
    def __init__(self, *, comparison: dict, font_name: str = PDF_FONT_NAME) -> None:
        super().__init__(width=500, height=245, font_name=font_name)
        self.comparison = comparison

    def draw(self) -> None:
        self.draw_frame("Динамика доходов и расходов")
        points = get_visible_line_points(self.comparison)
        if not points:
            self.draw_text(24, 115, "Нет данных за выбранный период", color="#64748B")
            return
        series = [
            ("Доходы", [float(point.get("income_rub") or 0) for point in points], "#22C55E"),
            ("Расходы", [float(point.get("expense_rub") or 0) for point in points], "#EF4444"),
        ]
        values = [value for _label, series_values, _color in series for value in series_values]
        max_value = max(values + [1])
        min_value = min(values + [0])
        value_range = max(max_value - min_value, 1)
        chart_x = 58
        chart_y = 48
        chart_width = 345
        chart_height = 152
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        for tick in range(5):
            y = chart_y + chart_height * tick / 4
            value = min_value + value_range * tick / 4
            canvas.line(chart_x, y, chart_x + chart_width, y)
            self.draw_right_text(chart_x - 6, y - 3, format_amount(value), size=6, color="#64748B")
        for label, series_values, color in series:
            coords = []
            for index, value in enumerate(series_values):
                x = chart_x + chart_width * index / max(len(series_values) - 1, 1)
                y = chart_y + chart_height * (value - min_value) / value_range
                coords.append((x, y))
            canvas.setStrokeColor(colors.HexColor(color))
            canvas.setLineWidth(2)
            for start, end in zip(coords, coords[1:]):
                canvas.line(start[0], start[1], end[0], end[1])
            canvas.setFillColor(colors.HexColor(color))
            for x, y in coords:
                canvas.circle(x, y, 2.4, stroke=0, fill=1)
        labels = [point.get("label", "") for point in points]
        for index, label in enumerate(labels[::2]):
            source_index = index * 2
            x = chart_x + chart_width * source_index / max(len(points) - 1, 1)
            self.draw_right_text(x + 10, chart_y - 14, truncate(label, 8), size=6, color="#475569")
        self.draw_legend(series)

    def draw_legend(self, series) -> None:
        x = 418
        y = 158
        for index, (label, _values, color) in enumerate(series):
            row_y = y - index * 18
            self.canv.setFillColor(colors.HexColor(color))
            self.canv.rect(x, row_y, 8, 8, stroke=0, fill=1)
            self.draw_text(x + 12, row_y, label, size=7)


def register_pdf_font() -> str:
    if PDF_FONT_NAME in pdfmetrics.getRegisteredFontNames():
        return PDF_FONT_NAME
    for font_path in PDF_FONT_CANDIDATES:
        if Path(font_path).exists():
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, font_path))
            return PDF_FONT_NAME
    return "Helvetica"


def format_amount(value) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        return "0"
    if abs(number) >= 1_000_000:
        return f"{number / 1_000_000:.1f}M"
    if abs(number) >= 1_000:
        return f"{number / 1_000:.0f}K"
    return f"{number:.0f}"


def truncate(value: str, limit: int) -> str:
    text = str(value or "")
    return text if len(text) <= limit else text[: limit - 1] + "…"
