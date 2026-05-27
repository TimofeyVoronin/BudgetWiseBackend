import csv
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from apps.finance.models import TransactionType


@dataclass(frozen=True)
class TransactionExportResult:
    content: bytes
    content_type: str
    filename: str


TRANSACTION_EXPORT_COLUMNS = [
    "Дата",
    "Тип",
    "Сумма",
    "Валюта",
    "Описание",
    "Категория",
    "Счёт",
    "Теги",
    "Дата создания",
    "Исходная сумма",
    "Исходная валюта",
]

TRANSACTION_EXPORT_FORMATS = {
    "csv",
    "xlsx",
    "pdf",
}

PDF_FONT_NAME = "BudgetWiseDejaVu"
PDF_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
    "/usr/local/share/fonts/DejaVuSans.ttf",
]


def build_transaction_export(
    *,
    transactions,
    export_format: str,
    currency_converter=None,
) -> TransactionExportResult:
    export_format = export_format.lower()
    rows = _build_export_rows(transactions, currency_converter=currency_converter)
    filename = _build_filename(export_format)

    if export_format == "csv":
        return TransactionExportResult(
            content=_build_csv(rows),
            content_type="text/csv; charset=utf-8",
            filename=filename,
        )

    if export_format == "xlsx":
        return TransactionExportResult(
            content=_build_xlsx(rows),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            filename=filename,
        )

    if export_format == "pdf":
        return TransactionExportResult(
            content=_build_pdf(rows),
            content_type="application/pdf",
            filename=filename,
        )

    raise ValueError(f"Unsupported export format: {export_format}")


def _build_export_rows(transactions, *, currency_converter=None) -> list[dict[str, str]]:
    rows = []

    for transaction in transactions:
        source_currency = _get_transaction_source_currency(transaction)
        source_amount = _get_signed_amount_value(transaction)
        display_amount, display_currency = _get_display_money(
            source_amount,
            source_currency=source_currency,
            currency_converter=currency_converter,
        )

        rows.append(
            {
                "Дата": transaction.operation_date.isoformat(),
                "Тип": _get_transaction_type_label(transaction.type),
                "Сумма": _format_money_amount(display_amount),
                "Валюта": display_currency,
                "Описание": transaction.description,
                "Категория": transaction.category.name,
                "Счёт": transaction.account.name,
                "Теги": _get_transaction_tags_label(transaction),
                "Дата создания": timezone.localtime(transaction.created_at).strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
                "Исходная сумма": _format_money_amount(source_amount),
                "Исходная валюта": source_currency,
            }
        )

    return rows



def _get_transaction_tags_label(transaction) -> str:
    tags_manager = getattr(transaction, "tags", None)

    if tags_manager is None:
        return ""

    return ", ".join(tag.name for tag in tags_manager.all())

def _get_transaction_type_label(transaction_type: str) -> str:
    if transaction_type == TransactionType.INCOME:
        return "Доход"

    if transaction_type == TransactionType.EXPENSE:
        return "Расход"

    return transaction_type


def _get_signed_amount_value(transaction) -> Decimal:
    amount = abs(transaction.amount)

    if transaction.type == TransactionType.EXPENSE:
        amount = -amount

    return amount


def _get_display_money(amount, *, source_currency: str, currency_converter=None) -> tuple[Decimal, str]:
    if currency_converter is None:
        return Decimal(amount), source_currency

    money = currency_converter.convert_to_display(
        amount,
        source_currency=source_currency,
    )
    return money.amount, money.currency


def _get_transaction_source_currency(transaction) -> str:
    account = getattr(transaction, "account", None)
    return getattr(account, "currency", None) or "RUB"


def _format_money_amount(value) -> str:
    return str(Decimal(value).quantize(Decimal("0.01")))


def _build_filename(export_format: str) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"transactions_{timestamp}.{export_format}"


def _build_csv(rows: list[dict[str, str]]) -> bytes:
    output = StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=TRANSACTION_EXPORT_COLUMNS,
        extrasaction="ignore",
    )

    writer.writeheader()
    writer.writerows(rows)

    return output.getvalue().encode("utf-8-sig")


def _build_xlsx(rows: list[dict[str, str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Операции"

    worksheet.append(TRANSACTION_EXPORT_COLUMNS)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="E5E7EB",
    )

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        worksheet.append([row[column] for column in TRANSACTION_EXPORT_COLUMNS])

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_pdf(rows: list[dict[str, str]]) -> bytes:
    font_name = _register_pdf_font()

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BudgetWiseTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=16,
        leading=20,
        spaceAfter=8,
    )
    cell_style = ParagraphStyle(
        "BudgetWiseCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8,
        leading=10,
    )

    story = [
        Paragraph("Экспорт финансовых операций", title_style),
    ]

    table_data = [
        [Paragraph(column, cell_style) for column in TRANSACTION_EXPORT_COLUMNS],
    ]

    for row in rows:
        table_data.append(
            [
                Paragraph(str(row[column]), cell_style)
                for column in TRANSACTION_EXPORT_COLUMNS
            ]
        )

    table = Table(
        table_data,
        repeatRows=1,
        colWidths=[
            20 * mm,
            15 * mm,
            18 * mm,
            15 * mm,
            44 * mm,
            27 * mm,
            27 * mm,
            33 * mm,
            28 * mm,
            25 * mm,
            20 * mm,
        ],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    story.append(table)
    document.build(story)

    return output.getvalue()


def _register_pdf_font() -> str:
    for font_path in PDF_FONT_CANDIDATES:
        path = Path(font_path)

        if path.exists():
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(path)))
            return PDF_FONT_NAME

    return "Helvetica"