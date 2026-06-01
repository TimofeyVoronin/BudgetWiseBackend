from __future__ import annotations

import base64
import csv
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle


@dataclass(frozen=True)
class FinancialCalendarExportResult:
    content: bytes
    content_type: str
    filename: str


FINANCIAL_CALENDAR_EXPORT_FORMAT_CSV = "csv"
FINANCIAL_CALENDAR_EXPORT_FORMAT_PDF = "pdf"
FINANCIAL_CALENDAR_EXPORT_FORMAT_XLSX = "xlsx"

FINANCIAL_CALENDAR_EXPORT_FORMATS = {
    FINANCIAL_CALENDAR_EXPORT_FORMAT_CSV,
    FINANCIAL_CALENDAR_EXPORT_FORMAT_PDF,
    FINANCIAL_CALENDAR_EXPORT_FORMAT_XLSX,
}

DEFAULT_FINANCIAL_CALENDAR_EXPORT_COLUMNS = {
    "actual": True,
    "balance": True,
    "events": True,
    "risks": True,
}

FINANCIAL_CALENDAR_EXPORT_COLUMN_TITLES = {
    "date": "Дата",
    "actual": "Факт ₽",
    "balance": "Баланс ₽",
    "events": "События",
    "risks": "Риск",
}

FINANCIAL_CALENDAR_RISK_LABELS = {
    "safe": "Безопасно",
    "caution": "Внимание",
    "risk": "Риск",
}

PDF_FONT_NAME = "BudgetWiseDejaVu"
PDF_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf",
    "/usr/local/share/fonts/DejaVuSans.ttf",
]


def normalize_financial_calendar_export_format(value: str | None) -> str:
    return str(value or FINANCIAL_CALENDAR_EXPORT_FORMAT_CSV).strip().lower()


def normalize_financial_calendar_export_columns(value: dict | None) -> dict[str, bool]:
    normalized = DEFAULT_FINANCIAL_CALENDAR_EXPORT_COLUMNS.copy()

    if not isinstance(value, dict):
        return normalized

    for key in normalized:
        raw_value = value.get(key)
        if raw_value is None:
            continue
        normalized[key] = _to_bool(raw_value)

    return normalized


def build_financial_calendar_export_file(
    *,
    rows: list[dict],
    export_format: str,
    columns: dict[str, bool] | None = None,
    year: int,
    month: int,
    currency_code: str = "RUB",
) -> FinancialCalendarExportResult:
    export_format = normalize_financial_calendar_export_format(export_format)
    columns = normalize_financial_calendar_export_columns(columns)
    selected_columns = _get_selected_columns(columns)
    column_titles = _get_export_column_titles(currency_code)
    table_rows = _build_table_rows(rows, selected_columns)
    filename = _build_filename(export_format, year, month)

    if export_format == FINANCIAL_CALENDAR_EXPORT_FORMAT_CSV:
        return FinancialCalendarExportResult(
            content=_build_csv(table_rows, selected_columns, column_titles=column_titles),
            content_type="text/csv;charset=utf-8",
            filename=filename,
        )

    if export_format == FINANCIAL_CALENDAR_EXPORT_FORMAT_XLSX:
        return FinancialCalendarExportResult(
            content=_build_xlsx(table_rows, selected_columns, column_titles=column_titles),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            filename=filename,
        )

    if export_format == FINANCIAL_CALENDAR_EXPORT_FORMAT_PDF:
        return FinancialCalendarExportResult(
            content=_build_pdf(table_rows, selected_columns, year=year, month=month, column_titles=column_titles),
            content_type="application/pdf",
            filename=filename,
        )

    raise ValueError("Формат экспорта должен быть csv, pdf или xlsx.")


def build_financial_calendar_download_url(export_result: FinancialCalendarExportResult) -> str:
    encoded_content = base64.b64encode(export_result.content).decode("ascii")
    return f"data:{export_result.content_type};base64,{encoded_content}"


def _to_bool(value) -> bool:
    if isinstance(value, bool):
        return value

    if isinstance(value, int):
        return bool(value)

    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "on"}:
        return True
    if normalized in {"0", "false", "no", "n", "off"}:
        return False

    return bool(value)


def _get_selected_columns(columns: dict[str, bool]) -> list[str]:
    selected = ["date"]

    for key in ("actual", "balance", "events", "risks"):
        if columns.get(key):
            selected.append(key)

    return selected


def _get_export_column_titles(currency_code: str) -> dict[str, str]:
    currency_label = _get_export_currency_label(currency_code)
    titles = FINANCIAL_CALENDAR_EXPORT_COLUMN_TITLES.copy()
    titles["actual"] = f"Факт {currency_label}"
    titles["balance"] = f"Баланс {currency_label}"
    return titles


def _get_export_currency_label(currency_code: str) -> str:
    code = str(currency_code or "RUB").strip().upper()
    return {
        "RUB": "₽",
        "USD": "$",
        "EUR": "€",
        "KZT": "₸",
        "CNY": "¥",
        "GBP": "£",
        "BTC": "₿",
        "ETH": "Ξ",
    }.get(code, code)


def _build_table_rows(rows: list[dict], selected_columns: list[str]) -> list[dict[str, str]]:
    table_rows: list[dict[str, str]] = []

    for row in rows:
        table_row = {
            "date": str(row.get("date") or ""),
            "actual": _money_to_export_value(row.get("actualRub")),
            "balance": _money_to_export_value(row.get("forecastRub")),
            "events": str(row.get("eventsSummary") or "—"),
            "risks": FINANCIAL_CALENDAR_RISK_LABELS.get(str(row.get("riskLevel") or "safe"), "Безопасно"),
        }
        table_rows.append({key: table_row[key] for key in selected_columns})

    return table_rows


def _money_to_export_value(value) -> str:
    if value in (None, ""):
        return "0.00"

    return str(Decimal(str(value)).quantize(Decimal("0.01")))


def _build_filename(export_format: str, year: int, month: int) -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"financial_calendar_{year}_{month:02d}_{timestamp}.{export_format}"


def _build_csv(
    rows: list[dict[str, str]],
    selected_columns: list[str],
    *,
    column_titles: dict[str, str],
) -> bytes:
    output = StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=[column_titles[column] for column in selected_columns],
        extrasaction="ignore",
    )

    writer.writeheader()
    for row in rows:
        writer.writerow(
            {
                column_titles[column]: row[column]
                for column in selected_columns
            }
        )

    return output.getvalue().encode("utf-8-sig")


def _build_xlsx(
    rows: list[dict[str, str]],
    selected_columns: list[str],
    *,
    column_titles: dict[str, str],
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Календарь"

    headers = [column_titles[column] for column in selected_columns]
    worksheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="E5E7EB")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        worksheet.append([row[column] for column in selected_columns])

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_pdf(
    rows: list[dict[str, str]],
    selected_columns: list[str],
    *,
    year: int,
    month: int,
    column_titles: dict[str, str],
) -> bytes:
    font_name = _register_pdf_font()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BudgetWiseCalendarTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=16,
        leading=20,
        spaceAfter=8,
    )
    cell_style = ParagraphStyle(
        "BudgetWiseCalendarCell",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8,
        leading=10,
    )

    headers = [column_titles[column] for column in selected_columns]
    table_data = [[Paragraph(header, cell_style) for header in headers]]

    for row in rows:
        table_data.append(
            [Paragraph(str(row[column]), cell_style) for column in selected_columns]
        )

    story = [
        Paragraph(f"Экспорт финансового календаря: {month:02d}.{year}", title_style),
    ]

    table = Table(
        table_data,
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    document.build(story)

    return output.getvalue()


def _register_pdf_font() -> str:
    for font_path in PDF_FONT_CANDIDATES:
        path = Path(font_path)

        if path.exists():
            pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(path)))
            return PDF_FONT_NAME

    return "Helvetica"
